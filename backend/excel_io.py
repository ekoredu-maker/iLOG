"""엑셀 읽기/쓰기 공통 도구.

- 엑셀 파일을 행 목록으로 읽고, 날짜는 항상 'YYYY-MM-DD' 로 맞춘다.
- 업로드용 작성 서식(템플릿)을 만든다.
"""
from __future__ import annotations

import base64
import io
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="555555")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E7EEF7")
HEAD_FONT = Font(bold=True)
FONT_NAME = "맑은 고딕"

_DATE_PATTERNS = [
    re.compile(r"^(\d{4})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})일?\.?$"),
    re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$"),  # 월/일/년 (미국식)
]


def normalize_date(value) -> str:
    """여러 형태의 날짜를 'YYYY-MM-DD' 로. 날짜가 아니면 빈 문자열."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 < value < 80000:  # 엑셀 날짜 일련번호
            return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
        return ""
    s = re.sub(r"\s+", " ", str(value).strip())
    s = re.sub(r"\(.\)$", "", s).strip()  # '2026-03-02(월)' 처리
    m = _DATE_PATTERNS[0].match(s)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = _DATE_PATTERNS[1].match(s)
        if not m:
            return ""
        mo, d, y = map(int, m.groups())
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return ""


def _clean(v):
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def read_rows(data: bytes, filename: str = "") -> list[list]:
    """첫 번째 시트를 2차원 목록으로. xlsx/xlsm, (xlrd가 있으면) xls 지원."""
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        try:
            import xlrd  # type: ignore
        except ImportError as e:  # pragma: no cover
            raise ValueError("예전 .xls 형식입니다. 엑셀에서 '다른 이름으로 저장 → .xlsx' 후 다시 올려주세요.") from e
        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
                else:
                    row.append(_clean(cell.value))
            rows.append(row)
        return _trim(rows)
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError("엑셀 파일을 열 수 없습니다. .xlsx 형식인지 확인해 주세요.") from e
    ws = wb.worksheets[0]
    rows = [[_clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return _trim(rows)


def _trim(rows: list[list]) -> list[list]:
    out = []
    for r in rows:
        r = list(r)
        while r and (r[-1] is None or r[-1] == ""):
            r.pop()
        out.append(["" if v is None else v for v in r])
    while out and not out[-1]:
        out.pop()
    return out


def read_rows_b64(b64: str, filename: str = "") -> list[list]:
    if "," in b64[:100]:
        b64 = b64.split(",", 1)[1]
    return read_rows(base64.b64decode(b64), filename)


# ----------------------------------------------------------------- 쓰기 도구
def style_table(ws, header_row: int, first_col: int, last_col: int, last_row: int, widths=None):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for r in range(header_row + 1, last_row + 1):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = BORDER
            if cell.alignment.horizontal is None:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(first_col + i)].width = w


def setup_print(ws, landscape=False, title_rows: str | None = None, header_text: str = ""):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    if header_text:
        ws.oddHeader.center.text = header_text
        ws.oddHeader.center.size = 9
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 9


def workbook_bytes(wb: Workbook) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def simple_template(title: str, headers: list[str], examples: list[list], widths=None, note: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for ex in examples:
        ws.append(ex)
    style_table(ws, 1, 1, len(headers), 1 + len(examples), widths)
    ws.freeze_panes = "A2"
    if note:
        ws2 = wb.create_sheet("작성 안내")
        for line in note.strip().splitlines():
            ws2.append([line])
        ws2.column_dimensions["A"].width = 100
    return workbook_bytes(wb)


TEMPLATES = {
    "subject_plan": dict(
        filename="지도계획_작성서식.xlsx",
        title="서식",
        headers=["날짜", "교시", "지도 단원", "학습목표", "지도내용(학습주제)"],
        examples=[["2026-03-02", 1, "1. 우리 학교", "학교를 알아요", "학교 탐방하기"],
                  ["2026-03-03", 2, "1. 우리 학교", "친구와 인사해요", "친구 얼굴 그리기"]],
        widths=[13, 7, 22, 30, 40],
        note="""1행(제목줄)은 지우지 마세요.
날짜: 2026-03-02, 2026.3.2, 엑셀 날짜 형식 모두 가능합니다.
교시: 그날 해당 과목이 들어있는 교시 번호입니다. (연간 시간표를 먼저 생성해야 반영됩니다)""",
    ),
    "students": dict(
        filename="학생명단_작성서식.xlsx",
        title="학생명단",
        headers=["번호", "이름", "성별", "전화번호", "비고"],
        examples=[[1, "김하늘", "여", "010-0000-0000", ""], [2, "이바다", "남", "", ""]],
        widths=[7, 12, 7, 18, 30],
        note="같은 번호가 이미 있으면 새로 추가하지 않고 정보를 갱신합니다.",
    ),
    "school_events": dict(
        filename="학사일정_작성서식.xlsx",
        title="학사일정",
        headers=["날짜", "행사명", "휴업"],
        examples=[["2026-03-02", "입학식·시업식", ""], ["2026-07-20~2026-08-16", "여름방학", "O"]],
        widths=[26, 30, 8],
        note="""기간은 '시작일~종료일' 로 적습니다.
휴업일(수업 없음)이면 휴업 칸에 O 를 적습니다.""",
    ),
    "curriculum_pack": dict(
        filename="출판사별_지도계획_작성서식.xlsx",
        title="지도계획",
        headers=["학년", "과목", "출판사", "학기", "차시", "단원", "학습목표", "지도내용", "교육과정"],
        examples=[[3, "수학", "예시출판사", 1, 1, "1. 덧셈과 뺄셈", "세 자리 수의 덧셈을 할 수 있다.", "받아올림이 없는 덧셈", "2022 개정"],
                  [3, "수학", "예시출판사", 1, "2~3", "1. 덧셈과 뺄셈", "받아올림이 있는 덧셈을 할 수 있다.", "받아올림이 한 번 있는 덧셈", "2022 개정"]],
        widths=[6, 10, 14, 6, 8, 22, 34, 34, 11],
        note="""한 파일에 여러 학년·과목·출판사를 함께 적을 수 있습니다.
과목 이름은 프로그램의 과목명(예: 국어, 수학, 바른생활)과 같아야 자동 연결됩니다.
학기: 1 또는 2. 비워두면 학기 구분 없이 1학기부터 차례로 배치합니다.
차시: '2~3' 처럼 적으면 두 차시에 같은 내용이 들어갑니다.""",
    ),
}


def make_template(kind: str) -> tuple[str, bytes]:
    t = TEMPLATES[kind]
    return t["filename"], simple_template(t["title"], t["headers"], t["examples"], t.get("widths"), t.get("note", ""))
