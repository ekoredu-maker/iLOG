"""학급경영록 출력 확장.

기존 reports.py의 안정된 출력 흐름을 유지하면서
- 학급교육과정 종합 기록
- 국가 학년군 기준 + 학교 편성 + 연간 계획 + 현재 이수 시수
- 과목설정 ↔ 지도계획 ↔ 지도내용 ↔ 평가계획 연동 현황
- 주간·월간 과목별 시수
- 수행평가 계획
을 학급경영록에 추가한다.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

from . import instruction_hours as ih
from . import reports
from .db import Database
from .excel_io import setup_print, workbook_bytes


RECORD_FIELDS = [
    ("classVision", "학급 교육 비전"),
    ("classGoals", "학급 교육 목표"),
    ("focus", "학급 운영 중점"),
    ("studentProfile", "학생·학급 실태 및 교육적 요구"),
    ("curriculumPrinciples", "교육과정 재구성·운영 원칙"),
    ("creativeActivities", "창의적 체험활동 운영"),
    ("schoolAutonomy", "학교자율시간 운영"),
    ("crossCurricular", "범교과·안전·인성 등 연계 계획"),
    ("assessmentPolicy", "교수·학습 및 평가 운영 방침"),
    ("reflection", "학기·학년도 운영 성찰"),
    ("changes", "교육과정 변경·보완 기록"),
]


def _assessment_rows(db: Database) -> list[list]:
    plans = sorted(db.get_all("assessment_plans"), key=lambda p: (p.get("date") or "9999-99-99", p.get("subjectName") or ""))
    rows = []
    for p in plans:
        tasks = db.query("assessment_tasks", "planId", p.get("planId"))
        rubrics = sorted(db.query("assessment_rubrics", "planId", p.get("planId")), key=lambda r: int(r.get("order") or 0))
        task = tasks[0] if tasks else {}
        rubric_text = "\n".join(f"{r.get('level') or ''}: {r.get('descriptor') or ''}" for r in rubrics if r.get("level") or r.get("descriptor"))
        term = p.get("term")
        rows.append([
            p.get("subjectName") or p.get("subjectShort") or "",
            f"{term}학기" if str(term) in ("1", "2") else "",
            p.get("unit") or "",
            p.get("title") or "",
            p.get("date") or "",
            p.get("domain") or "",
            p.get("method") or "",
            p.get("standard") or "",
            p.get("objective") or "",
            task.get("title") or "",
            task.get("description") or "",
            rubric_text,
        ])
    return rows


def _insert_curriculum_overview_sheet(wb, db: Database, rd: reports.ReportData) -> None:
    """학급교육과정의 서술형 기록과 교과 연동 현황을 맨 앞부분에 출력한다."""
    snap = ih.snapshot(db)
    rec = snap["classRecord"]
    try:
        idx = wb.sheetnames.index("기초시간표")
    except ValueError:
        idx = 1
    ws = wb.create_sheet("학급교육과정개요", idx)
    reports._title(ws, "학급교육과정 종합관리", rd.header_line, 9)

    row = 4
    for key, label in RECORD_FIELDS:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=rec.get(key) or "-")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        ws.cell(row=row, column=2).alignment = reports.Alignment(vertical="top", wrap_text=True) if hasattr(reports, "Alignment") else ws.cell(row=row, column=2).alignment
        ws.row_dimensions[row].height = 34 if rec.get(key) else 22
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="과목별 연동 현황")
    row += 1
    linked = []
    for r in snap["subjects"]:
        if r["scheduled"] == 0:
            status = "시간표 미배치"
        elif r["contentLessons"] == 0:
            status = "지도내용 미입력"
        elif r["assessmentPlans"] == 0 and r["evalPlans"] == 0:
            status = "평가계획 미연결"
        else:
            status = "연동"
        linked.append([
            r["name"],
            r["schoolPlan"] if r["schoolPlan"] is not None else "-",
            r["scheduled"],
            r["contentLessons"],
            f"{r['contentCoverage']}%",
            r["completed"],
            r["evalPlans"],
            r["assessmentPlans"],
            status,
        ])
    reports._table(
        ws, row,
        ["과목", "학교 편성", "연간 계획", "지도내용 입력", "내용 입력률", "현재 이수", "평가계획", "수행평가", "연동 상태"],
        linked or [["등록된 과목 없음", "-", 0, 0, "0%", 0, 0, 0, "-"]],
        [18, 12, 12, 14, 12, 12, 11, 11, 16],
        center_cols=range(2, 10),
    )
    ws.freeze_panes = f"A{row + 1}"
    setup_print(ws, landscape=True, header_text="학급교육과정 종합관리")


def _insert_hours_sheets(wb, db: Database, rd: reports.ReportData) -> None:
    # 기존 reports.py의 '시수현황'은 subject.hours 값을 연간 기준처럼 사용하므로 제거한다.
    if "시수현황" in wb.sheetnames:
        wb.remove(wb["시수현황"])

    try:
        idx = wb.sheetnames.index("기초시간표")
    except ValueError:
        idx = 1

    snap = ih.snapshot(db)

    # 1) 과목별 편성·계획·이수 및 지도·평가 연동
    ws = wb.create_sheet("과목별편성시간", idx)
    reports._title(ws, "교육과정 편성·계획·이수 시간", rd.header_line, 11)
    ws.cell(row=3, column=1, value=f"국가 기준은 {snap['band']}학년군 2년간 기준 수업 시수이며, 당해 학년 편성시수는 학교 교육과정에 따라 입력합니다.")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=11)

    nrows = [[r["group"], r["nationalBand"], r["schoolPlan"] if r["schoolPlan"] is not None else "-",
              r["scheduled"], r["completed"]] for r in snap["national"]]
    end = reports._table(ws, 5,
        ["교과(군)", "국가 학년군 기준(2년)", "학교 당해학년 편성합", "연간시간표 계획합", "현재 이수합"],
        nrows, [22, 18, 18, 18, 16], center_cols=(2, 3, 4, 5))

    start = end + 3
    ws.cell(row=start - 1, column=1, value="당해 학년 과목별 편성·지도·평가·이수")
    rows = []
    for r in snap["subjects"]:
        diff = r["diff"] if r["diff"] is not None else "-"
        rows.append([
            r["name"], r["schoolPlan"] if r["schoolPlan"] is not None else "-", r["scheduled"],
            r["t1"], r["t2"], r["contentLessons"], f"{r['contentCoverage']}%",
            r["evalPlans"], r["assessmentPlans"], r["completed"], diff,
        ])
    reports._table(ws, start,
        ["과목", "학교 편성", "연간 계획", "1학기", "2학기", "지도내용 입력", "내용 입력률", "평가계획", "수행평가", "현재 이수", "편성 대비 계획"],
        rows or [["등록된 과목 없음", "-", 0, 0, 0, 0, "0%", 0, 0, 0, "-"]],
        [18, 13, 12, 9, 9, 13, 12, 10, 10, 11, 16], center_cols=range(2, 12))
    ws.freeze_panes = f"A{start + 1}"
    setup_print(ws, landscape=True, header_text="교육과정 편성·계획·이수 시간")

    # 2) 주간 과목별 시수
    idx += 1
    weekly = snap["weekly"]
    ws = wb.create_sheet("주간과목별시수", idx)
    headers = ["주", "기간"] + weekly["subjects"] + ["합계"]
    reports._title(ws, "주간 과목별 수업 시수", rd.header_line, len(headers))
    rows = [[r["label"], f"{r['start']}~{r['end']}", *r["values"], r["total"]] for r in weekly["rows"]]
    widths = [8, 23] + [10] * len(weekly["subjects"]) + [9]
    reports._table(ws, 4, headers, rows or [["", "계획 없음"] + [0] * (len(headers) - 2)], widths, center_cols=range(1, len(headers) + 1))
    ws.freeze_panes = "C5"
    setup_print(ws, landscape=True, title_rows="4:4", header_text="주간 과목별 수업 시수")

    # 3) 월간 과목별 시수
    idx += 1
    monthly = snap["monthly"]
    ws = wb.create_sheet("월간과목별시수", idx)
    headers = ["월"] + monthly["subjects"] + ["합계"]
    reports._title(ws, "월간 과목별 수업 시수", rd.header_line, len(headers))
    rows = [[r["label"], *r["values"], r["total"]] for r in monthly["rows"]]
    widths = [14] + [10] * len(monthly["subjects"]) + [9]
    reports._table(ws, 4, headers, rows or [["계획 없음"] + [0] * (len(headers) - 1)], widths, center_cols=range(1, len(headers) + 1))
    ws.freeze_panes = "B5"
    setup_print(ws, landscape=True, title_rows="4:4", header_text="월간 과목별 수업 시수")


def class_book_xlsx(db: Database, include_incidents: bool = False) -> bytes:
    base = reports.class_book_xlsx(db, include_incidents)
    wb = load_workbook(io.BytesIO(base))
    rd = reports.ReportData(db)

    _insert_curriculum_overview_sheet(wb, db, rd)
    _insert_hours_sheets(wb, db, rd)

    rows = _assessment_rows(db)
    try:
        idx = wb.sheetnames.index("지도계획") + 1
    except ValueError:
        idx = len(wb.worksheets)
    ws = wb.create_sheet("수행평가계획", idx)
    headers = ["과목", "학기", "단원", "평가명", "시기", "영역", "방법", "성취기준", "학습목표", "수행과제명", "수행과제", "평가기준"]
    reports._title(ws, "수행평가 계획", rd.header_line, len(headers))
    if not rows:
        rows = [["", "", "", "등록된 수행평가 계획 없음", "", "", "", "", "", "", "", ""]]
    reports._table(ws, 4, headers, rows,
                   [10, 7, 20, 24, 12, 10, 11, 34, 34, 22, 45, 48],
                   center_cols=(1, 2, 5, 6, 7))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{4 + max(len(rows), 1)}"
    setup_print(ws, landscape=True, title_rows="4:4", header_text="수행평가 계획")
    return workbook_bytes(wb)


def _curriculum_overview_html(db: Database) -> str:
    snap = ih.snapshot(db)
    rec = snap["classRecord"]
    out = ['<h2 class="page-break">학급교육과정 종합관리</h2>']
    rows = [[label, rec.get(key) or "-"] for key, label in RECORD_FIELDS]
    out.append(reports._html_table(["항목", "기록"], rows, ["24%", "76%"]))
    out.append('<h3>과목별 연동 현황</h3>')
    linked = []
    for r in snap["subjects"]:
        if r["scheduled"] == 0:
            status = "시간표 미배치"
        elif r["contentLessons"] == 0:
            status = "지도내용 미입력"
        elif r["assessmentPlans"] == 0 and r["evalPlans"] == 0:
            status = "평가계획 미연결"
        else:
            status = "연동"
        linked.append([
            r["name"], r["schoolPlan"] if r["schoolPlan"] is not None else "-", r["scheduled"],
            r["contentLessons"], f"{r['contentCoverage']}%", r["completed"], r["evalPlans"], r["assessmentPlans"], status,
        ])
    out.append(reports._html_table(
        ["과목", "학교 편성", "연간 계획", "지도내용", "입력률", "현재 이수", "평가계획", "수행평가", "연동 상태"],
        linked, center=range(1, 9), empty="등록된 과목 없음"))
    return "".join(out)


def _hours_html(db: Database) -> str:
    snap = ih.snapshot(db)
    out = [
        '<h2 class="page-break">2. 과목별 편성·계획·이수 시간</h2>',
        f'<div class="infobox">국가 기준은 <b>{reports.esc(snap["band"])}학년군</b> 2년간 기준 수업 시수입니다. '
        '학교의 당해 학년 편성시수와 iLOG 연간 시간표의 계획·지도내용·평가·이수 현황을 함께 표시합니다.</div>',
    ]
    out.append(reports._html_table(
        ["교과(군)", "국가 학년군 기준(2년)", "학교 당해학년 편성합", "연간 계획합", "현재 이수합"],
        [[r["group"], r["nationalBand"], r["schoolPlan"] if r["schoolPlan"] is not None else "-", r["scheduled"], r["completed"]] for r in snap["national"]],
        ["28%", "18%", "18%", "18%", "18%"], center=range(1, 5)))
    out.append('<h3>당해 학년 과목별 편성·지도·평가·이수</h3>')
    out.append(reports._html_table(
        ["과목", "학교 편성", "연간 계획", "지도내용", "입력률", "평가계획", "수행평가", "현재 이수", "증감"],
        [[r["name"], r["schoolPlan"] if r["schoolPlan"] is not None else "-", r["scheduled"], r["contentLessons"],
          f"{r['contentCoverage']}%", r["evalPlans"], r["assessmentPlans"], r["completed"],
          r["diff"] if r["diff"] is not None else "-"] for r in snap["subjects"]],
        center=range(1, 9), empty="등록된 과목 없음"))

    weekly = snap["weekly"]
    out.append('<h2 class="page-break">3. 주간 과목별 시수</h2>')
    out.append(reports._html_table(
        ["주", "기간", *weekly["subjects"], "합계"],
        [[r["label"], f"{r['start']}~{r['end']}", *r["values"], r["total"]] for r in weekly["rows"]],
        center=range(2, len(weekly["subjects"]) + 3), empty="연간 시간표가 없습니다."))

    monthly = snap["monthly"]
    out.append('<h2 class="page-break">4. 월간 과목별 시수</h2>')
    out.append(reports._html_table(
        ["월", *monthly["subjects"], "합계"],
        [[r["label"], *r["values"], r["total"]] for r in monthly["rows"]],
        center=range(1, len(monthly["subjects"]) + 2), empty="연간 시간표가 없습니다."))
    return "".join(out)


def _assessment_html(db: Database, section_no: int = 7) -> str:
    rows = _assessment_rows(db)
    if not rows:
        return f'<h2 class="page-break">{section_no}. 수행평가 계획</h2><p class="empty">등록된 수행평가 계획이 없습니다.</p>'
    out = [f'<h2 class="page-break">{section_no}. 수행평가 계획</h2>']
    for r in rows:
        out.append(f"<h3>[{reports.esc(r[0])}] {reports.esc(r[3])}</h3>")
        out.append(
            '<div class="infobox">'
            f'학기: {reports.esc(r[1] or "-")} · 단원: {reports.esc(r[2] or "-")} · 시기: {reports.esc(r[4] or "-")}<br>'
            f'영역: {reports.esc(r[5] or "-")} · 방법: {reports.esc(r[6] or "-")}<br>'
            f'성취기준: {reports.esc(r[7] or "-")}<br>학습목표: {reports.esc(r[8] or "-")}'
            '</div>'
        )
        out.append(reports._html_table(
            ["수행과제명", "수행과제", "평가기준"], [[r[9], r[10], r[11]]],
            ["20%", "42%", "38%"], empty="계획 없음"))
    return "".join(out)


def class_book_html(db: Database, include_incidents: bool = False) -> str:
    html = reports.class_book_html(db, include_incidents)

    # 기존 3번 '과목별 이수 시간 현황'은 subject.hours 기반의 옛 집계라 제거한다.
    html = re.sub(
        r'<h2>3\. 과목별 이수 시간 현황</h2>.*?(?=<h2 class="page-break">4\. 과목별 연간 지도 계획</h2>)',
        '', html, count=1, flags=re.S,
    )

    old_timetable = '<h2>2. 주간 기초 시간표</h2>'
    if old_timetable in html:
        replacement = _curriculum_overview_html(db) + _hours_html(db) + '<h2 class="page-break">5. 주간 기초 시간표</h2>'
        html = html.replace(old_timetable, replacement, 1)
    html = html.replace('<h2 class="page-break">4. 과목별 연간 지도 계획</h2>', '<h2 class="page-break">6. 과목별 연간 지도 계획</h2>', 1)

    marker = '<h2 class="page-break">5. 학생 평가 기록</h2>'
    if marker not in html:
        return html
    html = html.replace(marker, _assessment_html(db, 7) + '<h2 class="page-break">8. 학생 평가 기록</h2>', 1)
    html = html.replace('<h2 class="page-break">6. 학생 출결 상황</h2>', '<h2 class="page-break">9. 학생 출결 상황</h2>', 1)
    html = html.replace('<h2 class="page-break">7. 학생 상담 일지</h2>', '<h2 class="page-break">10. 학생 상담 일지</h2>', 1)
    html = html.replace('<h2>8. 교외체험학습 현황</h2>', '<h2>11. 교외체험학습 현황</h2>', 1)
    html = html.replace('<h2 class="page-break">9. 학생 사안 기록 (대외비)</h2>', '<h2 class="page-break">12. 학생 사안 기록 (대외비)</h2>', 1)
    return html
