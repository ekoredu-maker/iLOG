"""서식 출력 - 엑셀(xlsx) 과 인쇄용 HTML(A4).

엑셀: 용지 A4, 머리행 반복, 쪽 번호, 열 너비/줄바꿈/테두리 지정.
HTML: 시스템 브라우저(Edge 등)로 열어 [인쇄] 또는 [PDF로 저장].
"""
from __future__ import annotations

import html
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from . import attendance as att
from .db import Database
from .excel_io import FONT_NAME, setup_print, style_table, workbook_bytes

WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]
DAY_KEYS = ["mon", "tue", "wed", "thu", "fri"]


def esc(v) -> str:
    return html.escape("" if v is None else str(v)).replace("\n", "<br>")


def weekday(d: str) -> str:
    try:
        return WEEKDAYS[date.fromisoformat(d).weekday()]
    except ValueError:
        return ""


def _num(s):
    try:
        return int(s.get("number") or 0)
    except (TypeError, ValueError):
        return 0


class ReportData:
    """출력에 필요한 자료를 한 번에 모아 둔다."""

    def __init__(self, db: Database):
        self.settings = db.get("settings", "global") or {}
        self.students = sorted(db.get_all("students"), key=_num)
        self.student_by_id = {s["studentId"]: s for s in self.students}
        self.subjects = db.get_all("subjects")
        self.timetable = (db.get("timetable_weekly", "weekly") or {}).get("grid") or {}
        self.annual = sorted(db.get_all("annual_schedule"), key=lambda d: d.get("date", ""))
        self.events = sorted(db.get_all("school_events"), key=lambda e: e.get("startDate", ""))
        self.eval_plans = sorted(db.get_all("eval_plans"), key=lambda p: (p.get("date") or "", p.get("subjectName") or ""))
        self.eval_scores = db.get_all("eval_scores")
        self.attendance = db.get_all("attendance")
        self.counseling = sorted(db.get_all("counseling"), key=lambda c: c.get("date") or "")
        self.experiential = sorted(db.get_all("experiential"), key=lambda e: e.get("startDate") or "")
        self.incidents = sorted(db.get_all("incidents"), key=lambda e: e.get("date") or "")
        self.tasks = db.get_all("tasks")

    # --------------------------------------------------------------- 공통
    @property
    def class_label(self) -> str:
        s = self.settings
        g, c = s.get("grade") or "", s.get("classNo") or ""
        return f"{g}학년 {c}반" if g or c else ""

    @property
    def header_line(self) -> str:
        s = self.settings
        parts = [f"{s.get('schoolYear') or date.today().year}학년도", s.get("schoolName") or "", self.class_label]
        if s.get("teacherName"):
            parts.append(f"담임 {s['teacherName']}")
        return " ".join(p for p in parts if p)

    def term_of(self, d: str) -> int:
        s = self.settings
        if s.get("term1Start") and s.get("term1End") and s["term1Start"] <= d <= s["term1End"]:
            return 1
        if s.get("term2Start") and s.get("term2End") and s["term2Start"] <= d <= s["term2End"]:
            return 2
        return 0

    def student_name(self, sid) -> str:
        s = self.student_by_id.get(sid)
        return s.get("name", "") if s else "(삭제된 학생)"

    def student_no(self, sid):
        s = self.student_by_id.get(sid)
        return s.get("number", "") if s else ""

    # ---------------------------------------------------------- 지도계획
    def plan_rows(self, short_name: str) -> list[dict]:
        rows = []
        for d in self.annual:
            for t in sorted(d.get("subjects") or [], key=lambda x: x.get("period") or 0):
                if t.get("name") != short_name:
                    continue
                tags = t.get("crossTags") or []
                cross = (f"[{', '.join(tags)}] " if tags else "") + (t.get("crossNote") or "")
                rows.append({
                    "date": d["date"], "weekday": weekday(d["date"]), "period": t.get("period", ""),
                    "unit": t.get("unit", ""), "objective": t.get("objective", ""),
                    "content": t.get("content", ""), "cross": cross.strip(), "term": self.term_of(d["date"]),
                })
        for i, r in enumerate(rows, 1):
            r["seq"] = i
        return rows

    def hours_table(self) -> list[dict]:
        counts = defaultdict(lambda: [0, 0, 0])
        for d in self.annual:
            term = self.term_of(d["date"])
            for t in d.get("subjects") or []:
                counts[t.get("name")][term] += 1
        out = []
        for s in self.subjects:
            c = counts.get(s.get("shortName"), [0, 0, 0])
            target = int(s.get("hours") or 0)
            total = c[1] + c[2] + c[0]
            out.append({"name": s.get("name"), "short": s.get("shortName"), "target": target or "-",
                        "t1": c[1], "t2": c[2], "total": total, "diff": (total - target) if target else "-"})
        return out

    # -------------------------------------------------------------- 출결
    @property
    def periods(self) -> int:
        try:
            return min(max(int(self.settings.get("periodsPerDay") or 6), 4), 8)
        except (TypeError, ValueError):
            return 6

    @property
    def school_days(self) -> set:
        return {d["date"] for d in self.annual if d.get("subjects")}

    def attendance_counts(self, start: str = "", end: str = "") -> list[dict]:
        """학생별 나이스 기준 출결 집계."""
        per = defaultdict(list)
        for a in self.attendance:
            per[a.get("studentId")].append(a)
        days = self.school_days
        out = []
        for s in self.students:
            c = att.student_counts(per.get(s["studentId"], []), days, start, end)
            c.update(att.totals(c))
            c.update({"studentId": s["studentId"], "number": s.get("number"), "name": s.get("name")})
            out.append(c)
        return out

    def attendance_summary(self) -> list[dict]:
        return [{"studentId": c["studentId"], "number": c["number"], "name": c["name"],
                 "absent": c["결석"], "late": c["지각"], "early": c["조퇴"], "result": c["결과"],
                 "recognized": c["출석인정"], "unclassified": c["미분류"]} for c in self.attendance_counts()]

    def attendance_details(self, sid=None) -> list[dict]:
        """결석/지각/조퇴 기록을 학생·날짜별로 묶는다 (교시는 모아서 표시)."""
        grouped: dict[tuple, dict] = {}
        for a in self.attendance:
            if a.get("status") == "출석" or (sid and a.get("studentId") != sid):
                continue
            reason = att.reason_of(a)
            k = (a.get("date"), a.get("studentId"), a.get("status"), reason)
            label = "출석인정" if reason == att.RECOGNIZED else f"{a.get('status')}({reason})"
            g = grouped.setdefault(k, {"date": a.get("date"), "studentId": a.get("studentId"),
                                       "status": label, "periods": [], "notes": []})
            g["periods"].append(a.get("period"))
            note = (a.get("note") or "").strip()
            if note and note != "이전 교시 연동" and note not in g["notes"]:
                g["notes"].append(note)
        rows = sorted(grouped.values(), key=lambda g: (g["date"] or "", _num(self.student_by_id.get(g["studentId"], {}))))
        for r in rows:
            r["periods"] = ",".join(str(p) for p in sorted(p for p in r["periods"] if p is not None))
            r["note"] = " / ".join(r["notes"])
            r["number"] = self.student_no(r["studentId"])
            r["name"] = self.student_name(r["studentId"])
        return rows

    def score_of(self, plan_id, sid) -> dict:
        for s in self.eval_scores:
            if s.get("planId") == plan_id and s.get("studentId") == sid:
                return s
        return {}


# ======================================================================
# 엑셀 출력
# ======================================================================
def _title(ws, text: str, sub: str, ncols: int):
    ws.cell(row=1, column=1, value=text).font = Font(name=FONT_NAME, bold=True, size=15)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=1, value=sub).font = Font(name=FONT_NAME, size=10, color="555555")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="right")
    ws.row_dimensions[1].height = 26


def _table(ws, start_row: int, headers: list[str], rows: list[list], widths=None, center_cols=()):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    for r_i, row in enumerate(rows, start_row + 1):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="center" if c_i in center_cols else None)
    style_table(ws, start_row, 1, len(headers), start_row + len(rows), widths)
    return start_row + len(rows)


def _sheet_title(name: str, used: set) -> str:
    bad = '[]:*?/\\'
    base = "".join("_" if ch in bad else ch for ch in name)[:28] or "시트"
    t, i = base, 2
    while t in used:
        t, i = f"{base[:25]}({i})", i + 1
    used.add(t)
    return t


def subject_plan_xlsx(db: Database, short_names: list[str] | None = None) -> bytes:
    rd = ReportData(db)
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    subjects = [s for s in rd.subjects if not short_names or s.get("shortName") in short_names]
    for s in subjects:
        rows = rd.plan_rows(s.get("shortName"))
        if not rows and not short_names:
            continue
        ws = wb.create_sheet(_sheet_title(s.get("name") or s.get("shortName"), used))
        headers = ["차시", "날짜", "요일", "교시", "지도 단원", "학습목표", "지도내용", "범교과"]
        _title(ws, f"{s.get('name')} 연간 지도 계획", rd.header_line, len(headers))
        data = [[r["seq"], r["date"], r["weekday"], r["period"], r["unit"], r["objective"], r["content"], r["cross"]] for r in rows]
        if not data:
            data = [["", "", "", "", "계획 없음", "", "", ""]]
        _table(ws, 4, headers, data, [6, 12, 5, 5, 22, 32, 40, 16], center_cols=(1, 2, 3, 4))
        ws.freeze_panes = "A5"
        setup_print(ws, landscape=True, title_rows="4:4", header_text=f"{s.get('name')} 지도 계획")
    if not wb.worksheets:
        ws = wb.create_sheet("지도계획")
        ws["A1"] = "연간 시간표에 수업이 없습니다. [연간 시간표 생성]을 먼저 해 주세요."
    return workbook_bytes(wb)


def eval_xlsx(db: Database, plan_id: str) -> tuple[str, bytes]:
    rd = ReportData(db)
    plan = next((p for p in rd.eval_plans if p.get("planId") == plan_id), None)
    if not plan:
        raise ValueError("평가 계획을 찾을 수 없습니다.")
    wb = Workbook()
    ws = wb.active
    ws.title = "평가결과"
    _title(ws, f"[{plan.get('subjectName')}] {plan.get('title')} 평가 결과", rd.header_line, 4)
    info = [("시기", plan.get("date")), ("영역", plan.get("domain")), ("평가요소", plan.get("element")),
            ("성취기준", plan.get("standard")), ("평가방법", plan.get("method"))]
    r = 4
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT_NAME, bold=True)
        ws.cell(row=r, column=2, value=v or "-").alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        r += 1
    rows = []
    for s in rd.students:
        sc = rd.score_of(plan_id, s["studentId"])
        rows.append([s.get("number"), s.get("name"), sc.get("score", ""), sc.get("note", "")])
    head = r + 1
    _table(ws, head, ["번호", "이름", "평가", "관찰 내용"], rows, [7, 12, 8, 70], center_cols=(1, 2, 3))
    setup_print(ws, title_rows=f"{head}:{head}", header_text=f"{plan.get('subjectName')} {plan.get('title')}")
    name = f"평가결과_{plan.get('subjectName')}_{plan.get('title')}.xlsx"
    return name, workbook_bytes(wb)


def class_book_xlsx(db: Database, include_incidents: bool = False) -> bytes:
    rd = ReportData(db)
    wb = Workbook()
    used: set = set()

    ws = wb.active
    ws.title = _sheet_title("학생명렬표", used)
    _title(ws, "학생 명렬표", rd.header_line, 5)
    _table(ws, 4, ["번호", "이름", "성별", "보호자 연락처", "비고"],
           [[s.get("number"), s.get("name"), s.get("gender"), s.get("guardianPhone", ""), s.get("note", "")] for s in rd.students],
           [7, 12, 7, 18, 40], center_cols=(1, 2, 3))
    setup_print(ws, title_rows="4:4", header_text="학생 명렬표")

    ws = wb.create_sheet(_sheet_title("기초시간표", used))
    _title(ws, "주간 기초 시간표", rd.header_line, 6)
    rows = [[f"{p}교시"] + [(rd.timetable.get(f"{d}-{p}") or {}).get("name", "") for d in DAY_KEYS] for p in range(1, rd.periods + 1)]
    _table(ws, 4, ["교시", "월", "화", "수", "목", "금"], rows, [8, 12, 12, 12, 12, 12], center_cols=range(1, 7))
    for r in range(5, 11):
        ws.row_dimensions[r].height = 30
    setup_print(ws)

    ws = wb.create_sheet(_sheet_title("시수현황", used))
    _title(ws, "과목별 이수 시간 현황", rd.header_line, 6)
    rows = [[h["name"], h["target"], h["t1"], h["t2"], h["total"], h["diff"]] for h in rd.hours_table()]
    _table(ws, 4, ["과목", "기준 시수", "1학기", "2학기", "계", "증감"], rows, [18, 10, 9, 9, 9, 9], center_cols=range(2, 7))
    setup_print(ws, title_rows="4:4")

    ws = wb.create_sheet(_sheet_title("지도계획", used))
    headers = ["과목", "차시", "날짜", "요일", "교시", "지도 단원", "학습목표", "지도내용", "범교과"]
    _title(ws, "과목별 연간 지도 계획", rd.header_line, len(headers))
    rows = []
    for s in rd.subjects:
        for r in rd.plan_rows(s.get("shortName")):
            rows.append([s.get("name"), r["seq"], r["date"], r["weekday"], r["period"], r["unit"], r["objective"], r["content"], r["cross"]])
    _table(ws, 4, headers, rows, [12, 6, 12, 5, 5, 20, 30, 36, 14], center_cols=(1, 2, 3, 4, 5))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{4 + max(len(rows), 1)}"
    setup_print(ws, landscape=True, title_rows="4:4", header_text="과목별 연간 지도 계획")

    ws = wb.create_sheet(_sheet_title("평가기록", used))
    headers = ["과목", "평가명", "시기", "번호", "이름", "평가", "관찰 내용"]
    _title(ws, "학생 평가 기록", rd.header_line, len(headers))
    rows = []
    for p in rd.eval_plans:
        for s in rd.students:
            sc = rd.score_of(p.get("planId"), s["studentId"])
            if sc.get("score") or sc.get("note"):
                rows.append([p.get("subjectName"), p.get("title"), p.get("date"), s.get("number"), s.get("name"), sc.get("score", ""), sc.get("note", "")])
    _table(ws, 4, headers, rows, [10, 22, 12, 6, 10, 7, 50], center_cols=(1, 3, 4, 5, 6))
    ws.auto_filter.ref = f"A4:G{4 + max(len(rows), 1)}"
    setup_print(ws, landscape=True, title_rows="4:4")

    ws = wb.create_sheet(_sheet_title("출결현황", used))
    _title(ws, "학생별 출결 현황 (일수, 나이스 기준)", rd.header_line, 7)
    rows = [[a["number"], a["name"], a["absent"], a["late"], a["early"], a["result"], a["recognized"]] for a in rd.attendance_summary()]
    end = _table(ws, 4, ["번호", "이름", "결석", "지각", "조퇴", "결과", "출석인정"], rows, [7, 12, 8, 8, 8, 8, 9], center_cols=range(1, 8))
    detail = rd.attendance_details()
    start = end + 3
    ws.cell(row=start - 1, column=1, value="상세 기록").font = Font(name=FONT_NAME, bold=True, size=12)
    _table(ws, start, ["날짜", "번호", "이름", "구분", "교시", "비고"],
           [[d["date"], d["number"], d["name"], d["status"], d["periods"], d["note"]] for d in detail],
           center_cols=(1, 2, 3, 4, 5))
    ws.column_dimensions["F"].width = 40
    setup_print(ws, title_rows="4:4")

    ws = wb.create_sheet(_sheet_title("상담일지", used))
    _title(ws, "학생 상담 일지", rd.header_line, 5)
    rows = [[c.get("date"), rd.student_no(c.get("studentId")), rd.student_name(c.get("studentId")), c.get("type"), c.get("content")] for c in rd.counseling]
    _table(ws, 4, ["날짜", "번호", "이름", "유형", "상담 내용"], rows, [12, 6, 10, 8, 70], center_cols=(1, 2, 3, 4))
    setup_print(ws, title_rows="4:4")

    ws = wb.create_sheet(_sheet_title("교외체험학습", used))
    _title(ws, "교외체험학습 현황", rd.header_line, 7)
    rows = []
    for e in rd.experiential:
        docs = [n for k, n in (("app", "신청서"), ("report", "보고서")) if (e.get("docs") or {}).get(k)]
        rows.append([rd.student_no(e.get("studentId")), e.get("studentName"), f"{e.get('startDate', '')}~{e.get('endDate', '')}",
                     e.get("type"), e.get("reason", ""), ", ".join(docs), "승인" if e.get("status") == "approved" else "신청"])
    _table(ws, 4, ["번호", "이름", "기간", "종류", "사유", "제출서류", "상태"], rows, [6, 10, 24, 14, 34, 14, 7], center_cols=(1, 2, 3, 4, 6, 7))
    setup_print(ws, landscape=True, title_rows="4:4")

    if include_incidents:
        ws = wb.create_sheet(_sheet_title("사안기록(대외비)", used))
        _title(ws, "학생 사안 기록 (대외비)", rd.header_line, 5)
        rows = [[i.get("date"), rd.student_no(i.get("studentId")), rd.student_name(i.get("studentId")), i.get("content"), ", ".join(i.get("measures") or [])] for i in rd.incidents]
        _table(ws, 4, ["날짜", "번호", "이름", "내용", "조치"], rows, [12, 6, 10, 60, 18], center_cols=(1, 2, 3))
        setup_print(ws, title_rows="4:4", header_text="대외비")
    return workbook_bytes(wb)


# ======================================================================
# 인쇄용 HTML
# ======================================================================
PRINT_CSS = """
@page { size: A4 portrait; margin: 14mm 12mm 16mm; }
@page landscape { size: A4 landscape; }
* { box-sizing: border-box; }
body { font-family: 'Malgun Gothic', '맑은 고딕', 'Apple SD Gothic Neo', sans-serif; font-size: 10.5pt;
       color: #111; margin: 0; background: #eef0f3; }
.sheet { background: #fff; max-width: 210mm; margin: 12px auto; padding: 14mm 12mm; box-shadow: 0 1px 6px rgba(0,0,0,.15); }
.sheet.wide { max-width: 297mm; }
h1 { font-size: 20pt; text-align: center; margin: 0 0 6mm; }
h2 { font-size: 13.5pt; margin: 8mm 0 3mm; padding-bottom: 1.5mm; border-bottom: 2px solid #333; break-after: avoid; }
h3 { font-size: 11.5pt; margin: 5mm 0 2mm; break-after: avoid; }
.meta { text-align: right; color: #444; font-size: 9.5pt; margin-bottom: 4mm; }
table { width: calc(100% - 2px); border-collapse: collapse; margin: 0 1px 4mm; table-layout: fixed; }
thead { display: table-header-group; }
tr { break-inside: avoid; }
th, td { border: 1px solid #555; padding: 1.6mm 1.8mm; vertical-align: middle; word-break: keep-all; overflow-wrap: anywhere; }
th { background: #e8edf3; font-weight: 700; text-align: center; }
td.c { text-align: center; }
td.neg { color: #c00; } td.pos { color: #0645ad; }
.cover { border: 2px solid #222; padding: 18mm 10mm; text-align: center; margin: 30mm 0 10mm; }
.cover .y { font-size: 26pt; font-weight: 700; margin-bottom: 8mm; }
.cover .s { font-size: 16pt; margin: 3mm 0; }
.infobox { border: 1px solid #999; padding: 2mm 3mm; margin-bottom: 2mm; font-size: 9.5pt; background: #fafafa; }
.page-break { break-before: page; }
.empty { color: #888; text-align: center; }
.toolbar { position: sticky; top: 0; background: #2c3e50; color: #fff; padding: 8px 16px; display: flex; gap: 10px; align-items: center; z-index: 10; }
.toolbar button { font-size: 14px; padding: 6px 16px; border: 0; border-radius: 6px; cursor: pointer; background: #fff; color: #2c3e50; font-weight: 700; }
.toolbar span { font-size: 13px; opacity: .85; }
@media print {
  body { background: #fff; }
  .toolbar { display: none; }
  .sheet { box-shadow: none; margin: 0; padding: 0; max-width: none; }
}
"""


def _page(title: str, body: str, wide: bool = False) -> str:
    orient = "@page { size: A4 landscape; }" if wide else ""
    return f"""<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8"><title>{esc(title)}</title>
<style>{PRINT_CSS}{orient}</style></head><body>
<div class="toolbar"><button onclick="window.print()">인쇄 / PDF 저장</button>
<span>인쇄 창에서 대상 프린터를 'PDF로 저장'으로 고르면 PDF 파일이 됩니다.</span></div>
<div class="sheet{' wide' if wide else ''}">{body}</div></body></html>"""


def _html_table(headers, rows, widths=None, center=(), empty="기록 없음") -> str:
    cols = "".join(f'<col style="width:{w}">' if w else "<col>" for w in (widths or [None] * len(headers)))
    head = "".join(f"<th>{esc(h)}</th>" for h in headers)
    if not rows:
        body = f'<tr><td class="empty" colspan="{len(headers)}">{esc(empty)}</td></tr>'
    else:
        body = "".join(
            "<tr>" + "".join(f'<td class="c">{esc(v)}</td>' if i in center else f"<td>{esc(v)}</td>" for i, v in enumerate(r)) + "</tr>"
            for r in rows
        )
    return f"<table><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def subject_plan_html(db: Database, short_names: list[str] | None = None) -> str:
    rd = ReportData(db)
    parts = [f'<div class="meta">{esc(rd.header_line)} · 출력일 {date.today().isoformat()}</div>']
    first = True
    for s in rd.subjects:
        if short_names and s.get("shortName") not in short_names:
            continue
        rows = rd.plan_rows(s.get("shortName"))
        if not rows and not short_names:
            continue
        parts.append(f'<h2 class="{"" if first else "page-break"}">{esc(s.get("name"))} 연간 지도 계획 <small>({len(rows)}차시)</small></h2>')
        first = False
        parts.append(_html_table(
            ["차시", "날짜", "교시", "지도 단원", "학습목표", "지도내용", "범교과"],
            [[r["seq"], f'{r["date"]}({r["weekday"]})', r["period"], r["unit"], r["objective"], r["content"], r["cross"]] for r in rows],
            ["6%", "13%", "6%", "17%", "22%", "24%", "12%"], center=(0, 1, 2), empty="계획 없음"))
    if first:
        parts.append('<p class="empty">연간 시간표에 수업이 없습니다. [연간 시간표 생성]을 먼저 해 주세요.</p>')
    return _page("과목별 연간 지도 계획", "".join(parts), wide=True)


def student_report_html(db: Database, sid: str) -> str:
    rd = ReportData(db)
    s = rd.student_by_id.get(sid)
    if not s:
        raise ValueError("학생을 찾을 수 없습니다.")
    summ = next((a for a in rd.attendance_summary() if a["studentId"] == sid), {})
    body = [f'<h1>{esc(s.get("number"))}번 {esc(s.get("name"))} 종합 기록</h1>',
            f'<div class="meta">{esc(rd.header_line)} · 출력일 {date.today().isoformat()}</div>',
            "<h2>출결</h2>",
            _html_table(["결석(일)", "지각(일)", "조퇴(일)", "결과(일)", "출석인정(일)"],
                        [[summ.get("absent", 0), summ.get("late", 0), summ.get("early", 0), summ.get("result", 0), summ.get("recognized", 0)]], center=range(5)),
            _html_table(["날짜", "구분", "교시", "비고"],
                        [[d["date"], d["status"], d["periods"], d["note"]] for d in rd.attendance_details(sid)],
                        ["18%", "12%", "14%", "56%"], center=(0, 1, 2)),
            "<h2>평가</h2>"]
    rows = []
    for p in rd.eval_plans:
        sc = rd.score_of(p.get("planId"), sid)
        if sc.get("score") or sc.get("note"):
            rows.append([p.get("subjectName"), p.get("title"), sc.get("score", ""), sc.get("note", "")])
    body.append(_html_table(["과목", "평가명", "결과", "관찰 내용"], rows, ["12%", "28%", "8%", "52%"], center=(0, 2)))
    body.append("<h2>상담</h2>")
    body.append(_html_table(["날짜", "유형", "내용"],
                            [[c.get("date"), c.get("type"), c.get("content")] for c in rd.counseling if c.get("studentId") == sid],
                            ["15%", "10%", "75%"], center=(0, 1)))
    return _page(f"{s.get('name')} 종합 기록", "".join(body))


def class_book_html(db: Database, include_incidents: bool = False) -> str:
    rd = ReportData(db)
    st = rd.settings
    b = [f"""<div class="cover"><div class="y">{esc(st.get('schoolYear') or date.today().year)}학년도 학급 경영록</div>
<div class="s">{esc(st.get('schoolName') or '')}</div><div class="s">{esc(rd.class_label)}</div>
<div class="s">담임 {esc(st.get('teacherName') or '')}</div><p>출력일 {date.today().isoformat()}</p></div>"""]

    b.append('<h2 class="page-break">1. 학생 명렬표</h2>')
    b.append(_html_table(["번호", "이름", "성별", "보호자 연락처", "비고"],
                         [[s.get("number"), s.get("name"), s.get("gender"), s.get("guardianPhone", ""), s.get("note", "")] for s in rd.students],
                         ["9%", "15%", "9%", "22%", "45%"], center=(0, 1, 2, 3)))

    b.append("<h2>2. 주간 기초 시간표</h2>")
    b.append(_html_table(["교시", "월", "화", "수", "목", "금"],
                         [[f"{p}교시"] + [(rd.timetable.get(f"{d}-{p}") or {}).get("name", "") for d in DAY_KEYS] for p in range(1, rd.periods + 1)],
                         center=range(6)))

    b.append("<h2>3. 과목별 이수 시간 현황</h2>")
    hrows = rd.hours_table()
    b.append(_html_table(["과목", "기준 시수", "1학기", "2학기", "계", "증감"],
                         [[h["name"], h["target"], h["t1"], h["t2"], h["total"], (f"+{h['diff']}" if isinstance(h["diff"], int) and h["diff"] > 0 else h["diff"])] for h in hrows],
                         center=range(1, 6)))

    b.append('<h2 class="page-break">4. 과목별 연간 지도 계획</h2>')
    for s in rd.subjects:
        rows = rd.plan_rows(s.get("shortName"))
        if not rows:
            continue
        b.append(f"<h3>■ {esc(s.get('name'))} ({len(rows)}차시)</h3>")
        b.append(_html_table(["차시", "날짜", "지도 단원", "학습목표", "지도내용", "범교과"],
                             [[r["seq"], f'{r["date"]}({r["weekday"]})', r["unit"], r["objective"], r["content"], r["cross"]] for r in rows],
                             ["7%", "15%", "18%", "23%", "25%", "12%"], center=(0, 1)))

    b.append('<h2 class="page-break">5. 학생 평가 기록</h2>')
    if not rd.eval_plans:
        b.append('<p class="empty">등록된 평가가 없습니다.</p>')
    for p in rd.eval_plans:
        b.append(f"<h3>[{esc(p.get('subjectName'))}] {esc(p.get('title'))}</h3>")
        b.append(f'<div class="infobox">시기: {esc(p.get("date") or "-")} · 영역: {esc(p.get("domain") or "-")} · 평가요소: {esc(p.get("element") or "-")}<br>'
                 f'성취기준: {esc(p.get("standard") or "-")}<br>평가방법: {esc(p.get("method") or "-")}</div>')
        rows = []
        for s in rd.students:
            sc = rd.score_of(p.get("planId"), s["studentId"])
            if sc.get("score") or sc.get("note"):
                rows.append([s.get("number"), s.get("name"), sc.get("score", ""), sc.get("note", "")])
        b.append(_html_table(["번호", "이름", "결과", "관찰 내용"], rows, ["8%", "12%", "8%", "72%"], center=(0, 1, 2)))

    b.append('<h2 class="page-break">6. 학생 출결 상황</h2>')
    b.append(_html_table(["번호", "이름", "결석", "지각", "조퇴", "결과", "출석인정"],
                         [[a["number"], a["name"], a["absent"], a["late"], a["early"], a["result"], a["recognized"]] for a in rd.attendance_summary()],
                         center=range(7)))
    b.append('<p style="font-size:9pt;color:#555">단위: 일 · 나이스 기준(결석한 날의 지각·조퇴·결과는 세지 않음, 출석인정은 출석으로 처리)</p>')
    b.append("<h3>상세 기록</h3>")
    b.append(_html_table(["날짜", "번호", "이름", "구분", "교시", "비고"],
                         [[d["date"], d["number"], d["name"], d["status"], d["periods"], d["note"]] for d in rd.attendance_details()],
                         ["13%", "7%", "11%", "14%", "10%", "45%"], center=(0, 1, 2, 3, 4)))

    b.append('<h2 class="page-break">7. 학생 상담 일지</h2>')
    b.append(_html_table(["날짜", "번호", "이름", "유형", "상담 내용"],
                         [[c.get("date"), rd.student_no(c.get("studentId")), rd.student_name(c.get("studentId")), c.get("type"), c.get("content")] for c in rd.counseling],
                         ["13%", "7%", "11%", "9%", "60%"], center=(0, 1, 2, 3)))

    b.append("<h2>8. 교외체험학습 현황</h2>")
    rows = []
    for e in rd.experiential:
        docs = [n for k, n in (("app", "신청서"), ("report", "보고서")) if (e.get("docs") or {}).get(k)]
        rows.append([e.get("studentName"), f"{e.get('startDate', '')}~{e.get('endDate', '')}", e.get("type"), e.get("reason", ""),
                     ", ".join(docs), "승인" if e.get("status") == "approved" else "신청"])
    b.append(_html_table(["학생", "기간", "종류", "사유", "제출서류", "상태"], rows,
                         ["11%", "24%", "13%", "28%", "15%", "9%"], center=(0, 1, 2, 4, 5)))

    if include_incidents:
        b.append('<h2 class="page-break">9. 학생 사안 기록 (대외비)</h2>')
        b.append(_html_table(["날짜", "이름", "내용", "조치"],
                             [[i.get("date"), rd.student_name(i.get("studentId")), i.get("content"), ", ".join(i.get("measures") or [])] for i in rd.incidents],
                             ["13%", "11%", "56%", "20%"], center=(0, 1)))
    return _page("종합 학급 경영록", "".join(b))


def stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")


# ======================================================================
# 나이스 출결 통계 (월별 / 학기별 / 학년)
# ======================================================================
def _month_ranges(days: set) -> list[tuple[str, str, str]]:
    months = sorted({d[:7] for d in days})
    return [(f"{int(m[5:7])}월", f"{m}-01", f"{m}-31") for m in months]


def neis_attendance_xlsx(db: Database) -> bytes:
    rd = ReportData(db)
    st = rd.settings
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    periods = []
    if st.get("term1Start") and st.get("term1End"):
        periods.append(("1학기", st["term1Start"], st["term1End"]))
    if st.get("term2Start") and st.get("term2End"):
        periods.append(("2학기", st["term2Start"], st["term2End"]))
    periods.append(("학년 전체", "", ""))
    periods += _month_ranges(rd.school_days)
    reasons = att.REASONS
    for label, start, end in periods:
        ws = wb.create_sheet(_sheet_title(label, used))
        ncol = 3 + 4 * len(reasons) + 3
        rng = f" ({start} ~ {end})" if start and label.endswith("학기") else ""
        _title(ws, f"출결 통계 - {label}{rng}", rd.header_line, ncol)
        # 2단 머리행
        top, sub = ["번호", "이름", "수업일수"], ["", "", ""]
        for s_ in att.STATUSES:
            top += [s_] + [""] * (len(reasons) - 1)
            sub += reasons
        top += ["출석인정", "출석일수", "미분류"]
        sub += ["", "", ""]
        for i, v in enumerate(top, 1):
            ws.cell(row=4, column=i, value=v)
        for i, v in enumerate(sub, 1):
            ws.cell(row=5, column=i, value=v)
        for i in (1, 2, 3, ncol - 2, ncol - 1, ncol):
            ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
        for k in range(4):
            c0 = 4 + k * len(reasons)
            ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + len(reasons) - 1)
        data = []
        for c in rd.attendance_counts(start, end if not label.endswith("월") else end):
            row = [c["number"], c["name"], c["수업일수"]]
            for s_ in att.STATUSES:
                row += [c[f"{s_}_{r}"] or "" for r in reasons]
            row += [c["출석인정"] or "", c["출석일수"], c["미분류"] or ""]
            data.append(row)
        for r_i, row in enumerate(data, 6):
            for c_i, v in enumerate(row, 1):
                ws.cell(row=r_i, column=c_i, value=v).alignment = Alignment(horizontal="center", vertical="center")
        style_table(ws, 5, 1, ncol, 5 + len(data), [6, 10, 7] + [5.5] * (4 * len(reasons)) + [7, 7, 7])
        style_table(ws, 4, 1, ncol, 4)
        note_row = 7 + len(data)
        ws.cell(row=note_row, column=1, value="※ 나이스 기준: 결석한 날의 지각·조퇴·결과는 세지 않음. 출석인정(교외체험학습 등)은 출석일수에 포함. "
                                               "'미분류'는 사유(질병/미인정/기타)가 없는 예전 기록이므로 출결 화면에서 사유를 지정해 주세요.").font = Font(name=FONT_NAME, size=9, color="555555")
        ws.freeze_panes = "D6"
        setup_print(ws, landscape=True, title_rows="4:5", header_text=f"출결 통계 - {label}")
    # 상세
    ws = wb.create_sheet(_sheet_title("상세기록", used))
    _title(ws, "출결 상세 기록", rd.header_line, 6)
    detail = rd.attendance_details()
    _table(ws, 4, ["날짜", "번호", "이름", "구분", "교시", "비고"],
           [[d["date"], d["number"], d["name"], d["status"], d["periods"], d["note"]] for d in detail],
           [12, 6, 10, 14, 10, 40], center_cols=(1, 2, 3, 4, 5))
    ws.auto_filter.ref = f"A4:F{4 + max(len(detail), 1)}"
    setup_print(ws, title_rows="4:4")
    return workbook_bytes(wb)
