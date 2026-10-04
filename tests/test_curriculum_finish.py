import io

from hwpx.document import HwpxDocument
from openpyxl import load_workbook

from backend import class_book_assessment
from backend import curriculum_report
from backend import instruction_hours as ih
from tests.test_curriculum_upgrade import _db


def _all_sheet_text(ws):
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_subject_order_and_school_year_week_numbers_flow_to_outputs(tmp_path):
    db = _db(tmp_path)
    try:
        settings = db.get("settings", "global")
        db.put("settings", {
            "id": ih.subject_order_key(settings),
            "order": ["joy", "kor", "math", "bar", "wise"],
        })
        # 한 주를 건너뛴 실제 수업 주를 추가해도 제1주, 제2주로 연속 번호를 매긴다.
        db.put("annual_schedule", {"date": "2026-03-16", "subjects": [
            {"period": 1, "name": "국", "unit": "듣기·말하기", "objective": "말한다", "content": "발표"},
        ]})

        snap = ih.snapshot(db, "2026-10-04")
        assert [r["name"] for r in snap["subjects"]][:2] == ["즐거운생활", "국어"]
        assert [r["label"] for r in snap["weekly"]["rows"]] == ["제1주", "제2주"]
        assert snap["weekly"]["subjects"][0] == "즐거운생활"
    finally:
        db.close()


def test_class_book_excel_uses_one_teacher_subject_list(tmp_path):
    db = _db(tmp_path)
    try:
        raw = class_book_assessment.class_book_xlsx(db)
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        overview = _all_sheet_text(wb["학급교육과정개요"])
        hours = _all_sheet_text(wb["과목별편성시간"])
        assert "담임 개설과목" in overview
        assert "편성시수" in overview
        assert "담임 개설과목" in hours
        assert "편성시수" in hours
        assert "학교 편성" not in overview
        assert "학교 편성" not in hours
    finally:
        db.close()


def test_hwpx_is_full_class_management_record(tmp_path):
    db = _db(tmp_path)
    try:
        db.put("students", {
            "studentId": "s1", "number": 1, "name": "가상학생", "gender": "여",
            "guardianPhone": "010-0000-0000", "note": "테스트",
        })
        db.put("attendance", {
            "attendanceId": "a1", "studentId": "s1", "date": "2026-03-03",
            "period": 1, "status": "지각", "reason": "질병", "note": "가상 기록",
        })
        db.put("counseling", {
            "logId": "c1", "studentId": "s1", "date": "2026-03-10",
            "type": "학생상담", "content": "학교생활 적응 상담",
        })
        db.put("experiential", {
            "expId": "e1", "studentId": "s1", "startDate": "2026-04-01", "endDate": "2026-04-01",
            "type": "가족동반여행", "destination": "제천", "reason": "가족 체험", "status": "approved",
        })

        name, raw = curriculum_report.class_curriculum_hwpx(db)
        assert name.startswith("학급경영록_") and name.endswith(".hwpx")
        doc = HwpxDocument.open(io.BytesIO(raw))
        text = doc.text.plain()
        for expected in (
            "학급경영록", "학급교육과정 종합 운영·기록", "학생 명렬표",
            "담임 개설과목", "주간 과목별 시수", "월간 과목별 시수",
            "주간 기초시간표", "연간 지도계획", "교수·학습 및 평가 계획",
            "학생 평가 기록", "학생 출결 상황", "학생 상담 일지", "교외체험학습 현황",
        ):
            assert expected in text
    finally:
        db.close()
