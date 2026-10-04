import base64
import io

from openpyxl import load_workbook
from hwpx.document import HwpxDocument

from backend.api import Api
from backend.db import Database


def _api(tmp_path):
    db = Database(tmp_path)
    assert db.unlock("1234")
    api = Api(db)
    api.db_put("settings", {"id": "global", "schoolYear": "2026", "schoolName": "테스트초", "grade": "5", "classNo": "1", "teacherName": "홍길동"})
    api.db_put("assessment_plans", {
        "planId": "assess_out", "subjectName": "과학", "subjectShort": "과", "term": 1,
        "unit": "3. 날씨와 우리 생활", "date": "2026-06-08", "title": "날씨 자료 분석 수행평가",
        "domain": "과정·기능", "method": "실험·실습", "standard": "[6과06-01] 날씨 자료를 해석한다.",
        "objective": "날씨 자료를 분석하여 생활과의 관계를 설명할 수 있다.",
        "lessonRefs": [{"date": "2026-06-05", "period": 3, "lessonSeq": 9, "content": "날씨 자료 비교"}],
    })
    api.db_put("assessment_tasks", {
        "taskId": "task_out", "planId": "assess_out", "title": "날씨 자료 분석하기",
        "description": "날씨 자료를 분석하고 결과를 근거로 설명한다.", "evidence": "활동지, 발표, 관찰 기록",
    })
    api.db_put_many("assessment_rubrics", [
        {"rubricId": "r1", "planId": "assess_out", "level": "잘함", "descriptor": "근거를 들어 정확히 설명함", "order": 1},
        {"rubricId": "r2", "planId": "assess_out", "level": "보통", "descriptor": "주요 내용을 설명함", "order": 2},
        {"rubricId": "r3", "planId": "assess_out", "level": "노력요함", "descriptor": "안내를 받아 설명함", "order": 3},
    ])
    return db, api


def test_assessment_plan_hwpx_builds_valid_document(tmp_path):
    db, api = _api(tmp_path)
    try:
        out = api.build_file("assessment_plan_hwpx", {"planId": "assess_out"})
        raw = base64.b64decode(out["b64"])
        assert out["filename"].endswith(".hwpx")
        assert len(raw) > 1000
        doc = HwpxDocument.open(io.BytesIO(raw))
        text = doc.text.plain()
        assert "수행평가 계획" in text
        assert "날씨 자료 분석 수행평가" in text
        assert "근거를 들어 정확히 설명함" in text
    finally:
        db.close()


def test_assessment_plans_bundle_hwpx(tmp_path):
    db, api = _api(tmp_path)
    try:
        out = api.build_file("assessment_plans_hwpx", {"term": 1})
        doc = HwpxDocument.open(io.BytesIO(base64.b64decode(out["b64"])))
        assert "날씨 자료 분석 수행평가" in doc.text.plain()
        assert "1학기" in out["filename"]
    finally:
        db.close()


def test_class_book_includes_assessment_plan_in_xlsx_and_html(tmp_path):
    db, api = _api(tmp_path)
    try:
        x = api.build_file("class_book_xlsx", {})
        wb = load_workbook(io.BytesIO(base64.b64decode(x["b64"])), data_only=True)
        assert "과목별편성시간" in wb.sheetnames
        assert "주간과목별시수" in wb.sheetnames
        assert "월간과목별시수" in wb.sheetnames
        assert "수행평가계획" in wb.sheetnames
        ws = wb["수행평가계획"]
        values = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
        assert "날씨 자료 분석 수행평가" in values
        assert "잘함: 근거를 들어 정확히 설명함" in values

        h = api.build_html("class_book", {})["html"]
        assert "2. 과목별 편성·계획·이수 시간" in h
        assert "3. 주간 과목별 시수" in h
        assert "4. 월간 과목별 시수" in h
        assert "7. 수행평가 계획" in h
        assert "날씨 자료 분석 수행평가" in h
        assert "8. 학생 평가 기록" in h
    finally:
        db.close()
