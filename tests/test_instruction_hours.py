import io

from openpyxl import load_workbook

from backend.db import Database
from backend import instruction_hours as ih
from backend import class_book_assessment


def make_db(tmp_path):
    db = Database(tmp_path)
    assert db.unlock("1234")
    db.put("settings", {
        "id": "global", "schoolYear": 2026, "schoolName": "테스트초등학교",
        "grade": 5, "classNo": 1, "teacherName": "가상담임",
        "term1Start": "2026-03-03", "term1End": "2026-07-20",
        "term2Start": "2026-08-24", "term2End": "2027-01-08",
        "periodsPerDay": 6,
    })
    subjects = [
        {"subjectId": "s_kor", "name": "국어", "shortName": "국", "hours": 6},
        {"subjectId": "s_soc", "name": "사회", "shortName": "사", "hours": 3},
        {"subjectId": "s_mor", "name": "도덕", "shortName": "도", "hours": 1},
        {"subjectId": "s_math", "name": "수학", "shortName": "수", "hours": 4},
        {"subjectId": "s_sci", "name": "과학", "shortName": "과", "hours": 3},
        {"subjectId": "s_prac", "name": "실과", "shortName": "실", "hours": 2},
        {"subjectId": "s_pe", "name": "체육", "shortName": "체", "hours": 3},
        {"subjectId": "s_mus", "name": "음악", "shortName": "음", "hours": 2},
        {"subjectId": "s_art", "name": "미술", "shortName": "미", "hours": 2},
        {"subjectId": "s_eng", "name": "영어", "shortName": "영", "hours": 3},
    ]
    db.put_many("subjects", subjects)
    db.put("settings", {
        "id": "instruction_hours_2026_5", "schoolYear": 2026, "grade": 5,
        "schoolPlan": {"국어": 204, "사회": 102, "도덕": 34, "수학": 136,
                       "과학": 102, "실과": 68, "체육": 102, "음악": 68, "미술": 68, "영어": 102},
    })
    db.put("settings", {
        "id": "class_curriculum_2026_5_1", "classVision": "서로 배우며 성장하는 학급",
        "classGoals": "자기주도성과 협력 역량을 기른다.", "focus": "읽기·토론·프로젝트 중심 운영",
    })
    db.put("timetable_weekly", {"id": "weekly", "grid": {"mon-1": {"id": "s_kor", "name": "국"}}})
    db.put_many("annual_schedule", [
        {"date": "2026-03-03", "subjects": [
            {"period": 1, "name": "국", "unit": "1", "objective": "가", "content": "가"},
            {"period": 2, "name": "수", "unit": "1", "objective": "나", "content": "나"},
            {"period": 3, "name": "사", "unit": "1", "objective": "다", "content": "다"},
            {"period": 4, "name": "도", "unit": "1", "objective": "라", "content": "라"},
        ]},
        {"date": "2026-09-01", "subjects": [
            {"period": 1, "name": "국", "unit": "2", "objective": "가", "content": "가"},
            {"period": 2, "name": "과", "unit": "2", "objective": "나", "content": "나"},
            {"period": 3, "name": "실", "unit": "2", "objective": "다", "content": "다"},
        ]},
        {"date": "2026-11-02", "subjects": [
            {"period": 1, "name": "국", "unit": "3", "objective": "가", "content": "가"},
            {"period": 2, "name": "영"},
        ]},
    ])
    db.put("eval_plans", {"planId": "ep1", "subjectId": "s_kor", "subjectName": "국어", "title": "국어 평가"})
    db.put("assessment_plans", {"planId": "ap1", "subjectName": "과학", "subjectShort": "과", "title": "과학 수행평가"})
    return db


def test_national_grade_band_and_subject_hours(tmp_path):
    db = make_db(tmp_path)
    try:
        snap = ih.snapshot(db, "2026-10-04")
        assert snap["band"] == "5-6"
        refs = {r["group"]: r for r in snap["national"]}
        assert refs["국어"]["nationalBand"] == 408
        assert refs["과학/실과"]["nationalBand"] == 340
        assert refs["합계"]["nationalBand"] == 2176

        rows = {r["name"]: r for r in snap["subjects"]}
        # 과목설정에 등록된 과목은 연간시간표에 아직 없어도 모두 남는다.
        assert set(rows) == {"국어", "사회", "도덕", "수학", "과학", "실과", "체육", "음악", "미술", "영어"}
        assert rows["체육"]["scheduled"] == 0
        assert rows["미술"]["contentLessons"] == 0

        assert rows["국어"]["schoolPlan"] == 204
        assert rows["국어"]["scheduled"] == 3
        assert rows["국어"]["completed"] == 2
        assert rows["국어"]["t1"] == 1
        assert rows["국어"]["t2"] == 2
        assert rows["국어"]["contentLessons"] == 3
        assert rows["국어"]["evalPlans"] == 1
        assert rows["과학"]["assessmentPlans"] == 1
        assert rows["영어"]["contentLessons"] == 0
        # subject.hours(주당 참고값 6)를 연간 기준시수로 사용하지 않는다.
        assert rows["국어"]["diff"] == 3 - 204
        assert snap["classRecord"]["classVision"] == "서로 배우며 성장하는 학급"
    finally:
        db.close()


def test_grade1_integrated_subjects_count_even_with_name_spacing(tmp_path):
    db = Database(tmp_path / "g1")
    assert db.unlock("1234")
    try:
        db.put("settings", {
            "id": "global", "schoolYear": 2026, "grade": 1, "classNo": 1,
            "term1Start": "2026-03-03", "term1End": "2026-07-20",
        })
        db.put_many("subjects", [
            {"subjectId": "k", "name": "국어", "shortName": "국", "hours": 7},
            {"subjectId": "m", "name": "수학", "shortName": "수", "hours": 4},
            {"subjectId": "b", "name": "바른생활", "shortName": "바생", "hours": 2},
            {"subjectId": "s", "name": "슬기로운생활", "shortName": "슬생", "hours": 3},
            {"subjectId": "j", "name": "즐거운생활", "shortName": "즐생", "hours": 4},
        ])
        db.put("annual_schedule", {"date": "2026-03-03", "subjects": [
            {"period": 1, "name": "국", "content": "국어"},
            {"period": 2, "name": "수", "content": "수학"},
            {"period": 3, "subjectName": "바른 생활", "content": "학교생활"},
            {"period": 4, "name": "슬생", "content": "관찰"},
            {"period": 5, "subjectShort": "즐생", "content": "놀이"},
        ]})
        rows = {r["name"]: r for r in ih.subject_rows(db, "2026-03-04")}
        assert rows["국어"]["scheduled"] == 1
        assert rows["수학"]["scheduled"] == 1
        assert rows["바른생활"]["scheduled"] == 1
        assert rows["슬기로운생활"]["scheduled"] == 1
        assert rows["즐거운생활"]["scheduled"] == 1
    finally:
        db.close()


def test_weekly_monthly_hours_follow_annual_schedule(tmp_path):
    db = make_db(tmp_path)
    try:
        weekly = ih.weekly_rows(db)
        monthly = ih.monthly_rows(db)
        assert len(weekly["rows"]) == 3
        assert [r["key"] for r in monthly["rows"]] == ["2026-03", "2026-09", "2026-11"]
        # 모든 설정 과목을 열로 유지한다.
        assert len(monthly["subjects"]) == 10
        march = monthly["rows"][0]
        assert march["total"] == 4
    finally:
        db.close()


def test_class_book_places_curriculum_and_hours_before_timetable_and_plan(tmp_path):
    db = make_db(tmp_path)
    try:
        data = class_book_assessment.class_book_xlsx(db)
        wb = load_workbook(io.BytesIO(data))
        names = wb.sheetnames
        assert "시수현황" not in names
        for title in ("학급교육과정개요", "과목별편성시간", "주간과목별시수", "월간과목별시수", "기초시간표", "지도계획"):
            assert title in names
        assert names.index("학급교육과정개요") < names.index("과목별편성시간")
        assert names.index("과목별편성시간") < names.index("주간과목별시수")
        assert names.index("주간과목별시수") < names.index("월간과목별시수")
        assert names.index("월간과목별시수") < names.index("기초시간표")
        assert names.index("기초시간표") < names.index("지도계획")

        overview = wb["학급교육과정개요"]
        text = "\n".join(str(c.value or "") for row in overview.iter_rows() for c in row)
        assert "서로 배우며 성장하는 학급" in text
        assert "평가계획" in text
        assert "수행평가" in text

        html = class_book_assessment.class_book_html(db)
        assert "학급교육과정 종합관리" in html
        assert "2. 과목별 편성·계획·이수 시간" in html
        assert "3. 주간 과목별 시수" in html
        assert "4. 월간 과목별 시수" in html
        assert "5. 주간 기초 시간표" in html
        assert "6. 과목별 연간 지도 계획" in html
        assert "3. 과목별 이수 시간 현황" not in html
    finally:
        db.close()
