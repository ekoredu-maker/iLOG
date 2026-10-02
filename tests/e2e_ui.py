"""화면 전체 흐름 자동 점검 (브라우저 E2E).

    pip install playwright && python -m playwright install chromium
    python tests/e2e_ui.py [스크린샷 폴더]

임시 데이터 폴더로 개발 서버를 띄우고, 실제 화면을 조작하며 확인한다.
화면 오류(JS 예외·콘솔 오류)가 하나라도 있으면 실패로 끝난다.
"""
import asyncio
import base64
import json
import socket
import sys
import tempfile
import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.async_api import async_playwright

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import devserver  # noqa: E402

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
OUT.mkdir(parents=True, exist_ok=True)
NEW_PW = "teacher#1"
results: list[tuple[str, bool, str]] = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), str(detail)))
    print(("  OK  " if cond else "  FAIL") + f" {name}" + (f"  ({detail})" if detail else ""))


def xlsx(rows, name):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    p = OUT / name
    wb.save(p)
    return str(p)


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


async def toast(page):
    await page.wait_for_timeout(250)
    return (await page.text_content("#toast-msg") or "").strip()


async def main(port):
    url = f"http://127.0.0.1:{port}/index.html"
    errors = []
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(accept_downloads=True, viewport={"width": 1400, "height": 900}, locale="ko-KR", timezone_id="Asia/Seoul")
        page = await ctx.new_page()
        page.on("pageerror", lambda e: errors.append("pageerror: " + str(e)))
        page.on("console", lambda m: m.type == "error" and "501" not in m.text and errors.append("console: " + m.text))
        page.on("dialog", lambda d: asyncio.ensure_future(d.accept()))

        await page.goto(url)
        await page.wait_for_selector("#login-hint:has-text('1234')")
        await page.fill("#login-password", "0000")
        await page.click("#login-btn")
        await page.wait_for_selector("#login-msg:visible")
        check("틀린 비밀번호 거부", "맞지 않습니다" in await page.text_content("#login-msg"))
        await page.fill("#login-password", "1234")
        await page.click("#login-btn")
        await page.wait_for_selector("#firstPwModal.show")
        check("첫 로그인 시 비밀번호 변경 강제", True)
        await page.fill("#first-pw-new", "1234")
        await page.fill("#first-pw-new2", "1234")
        await page.click("#first-pw-btn")
        check("1234 재사용 거부", "1234" in await page.text_content("#first-pw-msg"))
        await page.fill("#first-pw-new", NEW_PW)
        await page.fill("#first-pw-new2", NEW_PW)
        await page.click("#first-pw-btn")
        await page.wait_for_selector("#recoveryCodeModal.show")
        code = (await page.text_content("#rc-code")).strip()
        check("복구 코드 발급", len(code.replace("-", "")) == 20, code)
        check("보관 확인 전 닫기 불가", await page.is_disabled("#rc-close-btn"))
        await page.check("#rc-ack")
        await page.click("#rc-close-btn")
        await page.wait_for_selector("#login-overlay", state="hidden")
        await page.screenshot(path=str(OUT / "01_dashboard.png"))

        # 설정
        await page.click('[data-target="settings"]')
        for k, v in {"schoolName": "제천테스트초등학교", "schoolYear": "2026", "grade": "3", "classNo": "2", "teacherName": "박주가리",
                     "term1Start": "2026-03-02", "term1End": "2026-07-17", "term2Start": "2026-08-17", "term2End": "2027-01-08",
                     "evalScale": "잘함, 보통, 노력요함"}.items():
            await page.fill(f"#{k}", v)
        await page.select_option("#periodsPerDay", "5")
        await page.click("#save-settings-btn")
        await page.wait_for_timeout(300)

        # 시간표 (5교시 반영)
        await page.click('[data-target="curriculum"]')
        await page.click("#tab-timetable")
        await page.wait_for_timeout(400)
        check("교시 수 설정 반영(5교시)", await page.locator("#timetable-grid tr").count() == 5)
        pal = page.locator("#draggable-subjects .draggable-item")
        idx = {n.strip(): i for i, n in enumerate(await pal.all_text_contents())}
        for d in ["mon", "tue", "wed", "thu", "fri"]:
            for per, key in {1: "국국어", 2: "수수학", 3: "사사회", 4: "과과학"}.items():
                await pal.nth(idx[key]).drag_to(page.locator(f"#{d}-{per}"))
        await page.click("#save-timetable-btn")
        await page.wait_for_timeout(300)

        # 학사일정 + 연간 시간표
        await page.click('[data-target="annual-plan"]')
        await page.wait_for_timeout(400)
        await page.set_input_files("#school-event-excel-file", xlsx([["날짜", "행사명", "휴업"], ["2026-05-05", "어린이날", "O"], ["2026-03-02", "입학식", ""]], "ev.xlsx"))
        await page.click("#school-event-excel-btn")
        check("학사일정 엑셀", "2건" in await toast(page))
        await page.click("#generate-annual-btn")
        await page.wait_for_timeout(1500)
        check("연간 시간표 생성", "일 생성" in await toast(page))

        # 지도계획 라이브러리
        rows = [["학년", "과목", "출판사", "학기", "차시", "단원", "학습목표", "지도내용", "교육과정"]]
        rows += [[3, "과학", "테스트출판", sem, i, f"{sem}학기 단원", f"목표{i}", f"과학 {sem}학기 {i}차시", "2022 개정"] for sem in (1, 2) for i in range(1, 71)]
        await page.click("#open-curriculum-btn")
        await page.wait_for_timeout(700)
        await page.set_input_files("#cur-file", xlsx(rows, "pack.xlsx"))
        await page.click("#cur-import-btn")
        check("지도계획 가져오기", "140차시" in await toast(page))
        sci = page.locator("#cur-subject-body tr", has_text="과학").first
        await sci.locator("select").select_option(label="테스트출판 · 2022 개정 (140차시)")
        await page.wait_for_timeout(700)
        await sci.locator("button").click()
        await page.wait_for_timeout(1200)
        check("지도계획 적용", "140칸" in await toast(page))
        await page.click("#curriculumModal .btn-close")
        await page.wait_for_timeout(400)

        # 학생
        await page.click('[data-target="students"]')
        await page.click('[data-bs-target="#batchStudentModal"]')
        await page.wait_for_timeout(400)
        await page.set_input_files("#batch-excel-file", xlsx([["번호", "이름", "성별"]] + [[i, f"학생{i:02d}", "여" if i % 2 else "남"] for i in range(1, 26)] + [[26, "O'Brien<b>", "남"]], "st.xlsx"))
        await page.click("#batch-excel-btn")
        await page.wait_for_timeout(800)
        check("학생 일괄 등록(특수문자 이름 포함)", (await page.text_content("#student-count")) == "26명")

        # 출결: 결석(미인정) → 정정
        await page.evaluate("openAttendanceModalForDate('2026-03-03', 2, '수')")
        await page.wait_for_timeout(700)
        sid = await page.evaluate("(async()=> (await StudentRepo.getAll())[0].studentId)()")
        sid2 = await page.evaluate("(async()=> (await StudentRepo.getAll())[1].studentId)()")
        check("출석이면 구분 선택 잠김", await page.is_disabled(f"#reason_{sid}"))
        await page.click(f'label[for="att_{sid}1"]')
        await page.select_option(f"#reason_{sid}", "미인정")
        await page.click(f'label[for="att_{sid2}4"]')  # 결과
        await page.select_option(f"#reason_{sid2}", "질병")
        await page.screenshot(path=str(OUT / "02_attendance.png"))
        await page.click("#save-attendance-btn")
        await page.wait_for_timeout(700)
        st = await page.evaluate(f"AttendanceRepo.getStudentStats('{sid}')")
        st2 = await page.evaluate(f"AttendanceRepo.getStudentStats('{sid2}')")
        check("결석(미인정) 전 교시 기록", st["absent"] == 1 and all(r.get("reason") == "미인정" for r in st["records"]), len(st["records"]))
        check("결과는 해당 교시만", st2["result"] == 1 and len(st2["records"]) == 1)

        # 체험학습 승인 → 출석인정
        await page.click('[data-target="experiential"]')
        await page.wait_for_timeout(400)
        await page.select_option("#exp-student", index=2)
        await page.fill("#exp-start", "2026-04-01")
        await page.fill("#exp-end", "2026-04-03")
        await page.check("#exp-approved")
        await page.click("#add-exp-btn")
        check("체험학습 승인 → 출석인정 반영", "출석인정 3일" in await toast(page))
        sid3 = await page.evaluate("(async()=> (await StudentRepo.getAll())[2].studentId)()")
        st3 = await page.evaluate(f"AttendanceRepo.getStudentStats('{sid3}')")
        check("출석인정은 결석으로 세지 않음", st3["absent"] == 0 and st3["recognized"] == 3)
        await page.locator("#exp-list-body .badge.cursor-pointer").first.click()
        await page.wait_for_timeout(700)
        st3 = await page.evaluate(f"AttendanceRepo.getStudentStats('{sid3}')")
        check("승인 취소 → 출석인정 기록 삭제", st3["recognized"] == 0)

        # 평가: 척도 설정 반영
        await page.click('[data-target="evaluation"]')
        await page.wait_for_timeout(300)
        await page.click('[data-bs-target="#addEvalModal"]')
        await page.wait_for_timeout(400)
        await page.fill("#eval-title", "관찰 평가")
        await page.click("#add-eval-plan-btn")
        await page.wait_for_timeout(500)
        await page.locator("#eval-plan-list .list-group-item").first.click()
        await page.wait_for_timeout(500)
        opts = await page.locator(f"#sv_{sid} option").all_text_contents()
        check("평가 척도 설정 반영", opts == ["-", "잘함", "보통", "노력요함"], opts)

        # 사안 (비밀번호 확인, 로그아웃 없음)
        await page.click('[data-target="incidents"]')
        await page.wait_for_selector("#passwordModal.show")
        await page.fill("#pw-modal-input", "1234")
        await page.click("#pw-modal-ok")
        try:
            await page.wait_for_selector("#pw-modal-msg:visible", timeout=3000)
            rejected = True
        except Exception:
            rejected = False
        check("사안: 옛 비밀번호 거부", rejected and await page.is_visible("#passwordModal.show"))
        await page.fill("#pw-modal-input", NEW_PW)
        await page.click("#pw-modal-ok")
        await page.wait_for_timeout(500)
        await page.fill("#incident-content", "쉬는 시간 다툼")
        await page.click("#add-incident-btn")
        await page.wait_for_timeout(500)
        check("사안 저장 후 로그인 유지", await page.is_hidden("#login-overlay") and await page.locator("#incident-list-body tr").count() == 1)

        # ---------------- 서식 출력
        from hwpx.document import HwpxDocument
        import io as _io
        sys.path.insert(0, str(ROOT))
        from backend.hwpx_builder import HwpxBuilder

        async def hwpx_text(click_sel):
            async with page.expect_download() as dl:
                await page.click(click_sel)
            d = await dl.value
            data = Path(await d.path()).read_bytes()
            return d.suggested_filename, HwpxDocument.open(_io.BytesIO(data)).text.plain()

        await page.click('[data-target="settings"]')
        await page.fill("#principalName", "홍교장")
        await page.fill("#expDomesticDays", "7")
        await page.click("#save-settings-btn")
        await page.wait_for_timeout(300)
        await page.click('[data-target="forms"]')
        await page.wait_for_timeout(600)
        await page.fill("#wk-date", "2026-03-04")
        await page.dispatch_event("#wk-date", "change")
        await page.wait_for_timeout(500)
        check("주간: 월요일로 맞춤", (await page.input_value("#wk-date")) == "2026-03-02")
        await page.fill('[data-wk="morning"][data-day="mon"]', "아침 독서")
        await page.fill('[data-wk="prep"][data-day="wed"]', "리코더")
        await page.fill("#wk-general", "금요일 학부모 공개수업")
        name, t = await hwpx_text("#wk-hwpx")
        check("주간학습안내 한글", "주간학습안내" in t and "리코더" in t and "금요일 학부모 공개수업" in t and "과학 1학기" in t, name)
        async with ctx.expect_page() as np:
            await page.click("#wk-print")
        rp = await np.value
        await rp.wait_for_load_state()
        await rp.pdf(path=str(OUT / "weekly.pdf"), prefer_css_page_size=True)
        await rp.close()

        await page.click("#tab-form-card")
        await page.wait_for_timeout(700)
        await page.fill(f'[data-card="{sid}"]', "친구를 배려하며 수업에 적극적으로 참여함.")
        name, t = await hwpx_text("#card-hwpx")
        check("가정통지표 한글(전체)", t.count("가정통지표") == 26 and "친구를 배려하며" in t, name)

        await page.click("#tab-form-absence")
        await page.wait_for_timeout(700)
        await page.select_option("#abs-student", sid)
        await page.dispatch_event("#abs-student", "change")
        await page.wait_for_timeout(500)
        opts = await page.locator("#abs-date option").all_text_contents()
        check("결석신고서: 날짜 목록", any("결석(미인정)" in o for o in opts), opts)
        name, t = await hwpx_text("#abs-hwpx")
        check("결석 미인정 → 학부모 확인서", "학부모 확인서" in t and "☑ 미인정" in t, name)

        await page.click("#tab-form-exp")
        await page.wait_for_timeout(700)
        await page.click("#exp-form-edit")
        await page.wait_for_selector("#editExpModal.show")
        await page.fill("#exp-edit-destination", "제주도")
        await page.fill("#exp-edit-guardianName", "김엄마")
        await page.fill("#exp-edit-plan", "1일차 성산일출봉")
        await page.click("#save-exp-edit-btn")
        await page.wait_for_timeout(800)
        name, t = await hwpx_text('[data-expform="exp_application"][data-fmt="hwpx"]')
        check("체험학습 신청서 한글", "제주도" in t and "김엄마" in t and "허용 7일" in t, name)

        tb = HwpxBuilder()
        tb.para("{{이름}} ({{학년반번호}})")
        tb.table([["성명", ""], ["목적지", ""]], [30, 140])
        (OUT / "school_tpl.hwpx").write_bytes(tb.to_bytes())
        await page.click("#tab-form-tpl")
        await page.wait_for_timeout(600)
        row = page.locator("#tpl-body tr", has_text="교외체험학습 신청서")
        await row.locator("input[type=file]").set_input_files(str(OUT / "school_tpl.hwpx"))
        await row.locator("button", has_text="등록").click()
        await page.wait_for_timeout(700)
        check("학교 양식 등록", "school_tpl.hwpx" in (await page.locator("#tpl-body").inner_text()))
        await page.click("#tab-form-exp")
        await page.wait_for_timeout(600)
        check("학교 양식 사용 기본 선택", await page.is_checked("#exp-use-tpl"))
        name, t = await hwpx_text('[data-expform="exp_application"][data-fmt="hwpx"]')
        check("학교 양식 채우기", "목적지\t제주도" in t and "학생03" in t and "학교장허가" not in t, t[:80])
        await page.screenshot(path=str(OUT / "03_forms.png"))

        # 출력
        await page.click('[data-target="stats"]')
        await page.wait_for_timeout(300)
        async with page.expect_download() as dl:
            await page.click("#neis-att-btn")
        d = await dl.value
        await d.save_as(OUT / "neis.xlsx")
        wb = load_workbook(OUT / "neis.xlsx")
        ws = wb["1학기"]
        row = next(r for r in ws.iter_rows(min_row=6, values_only=True) if r[0] == 1)
        check("나이스 출결 통계(1번 결석·미인정 1일)", row[4] == 1, row[:8])
        await page.click('[data-target="data-manage"]')
        await page.wait_for_timeout(400)
        async with page.expect_download() as dl:
            await page.click("#excel-all-report-btn")
        await (await dl.value).save_as(OUT / "book.xlsx")
        async with ctx.expect_page() as np:
            await page.click("#print-all-report-btn")
        rp = await np.value
        await rp.wait_for_load_state()
        await rp.pdf(path=str(OUT / "report.pdf"), format="A4")
        html = await rp.content()
        check("종합 리포트(이스케이프·사안 제외)", "O'Brien&lt;b&gt;" in html and "쉬는 시간 다툼" not in html)
        await rp.close()
        async with page.expect_download() as dl:
            await page.click("#backup-btn")
        await (await dl.value).save_as(OUT / "backup.json")
        backup = (OUT / "backup.json").read_text(encoding="utf-8")
        check("백업 파일 암호화", "학생01" not in backup and "ilog-backup-enc" in backup)

        # 스케줄러
        await page.click('[data-target="scheduler-modal-trigger"]')
        await page.wait_for_timeout(1200)
        check("달력 한국어", "년" in (await page.text_content(".fc-toolbar-title")))
        await page.click("#fullSchedulerModal .btn-close")
        await page.wait_for_timeout(400)

        # 잠금 → 복구 코드로 비밀번호 재설정
        await page.click("#lock-btn")
        await page.wait_for_selector("#login-overlay", state="visible")
        await page.reload()
        await page.wait_for_selector("#forgot-pw-link:visible")
        await page.click("#forgot-pw-link")
        await page.wait_for_selector("#recoverModal.show")
        await page.fill("#recover-code", code.lower())
        await page.fill("#recover-new", "after-recover")
        await page.fill("#recover-new2", "after-recover")
        await page.click("#recover-btn")
        await page.wait_for_selector("#recoveryCodeModal.show")
        await page.check("#rc-ack")
        await page.click("#rc-close-btn")
        await page.wait_for_selector("#login-overlay", state="hidden")
        n = await page.evaluate("(async()=> (await StudentRepo.getAll()).length)()")
        check("복구 코드로 재설정 후 데이터 유지", n == 26, n)

        # 암호화 백업 복구(같은 PC) + 예전 v9 백업 복구
        await page.click('[data-target="data-manage"]')
        await page.set_input_files("#importFile", str(OUT / "backup.json"))
        await page.click("#restore-btn")
        await page.wait_for_load_state()
        await page.wait_for_timeout(1500)
        await page.fill("#login-password", "after-recover")
        await page.click("#login-btn")
        await page.wait_for_selector("#login-overlay", state="hidden")
        check("암호화 백업 복구", await page.evaluate("(async()=> (await StudentRepo.getAll()).length)()") == 26)
        old = {"settings": [{"id": "global", "schoolName": "옛학교"}, {"id": "security", "password": "4321"}],
               "students": [{"studentId": "std_1", "number": 1, "name": "옛학생", "gender": "남"}],
               "attendance": [{"attendanceId": "a1", "date": "2026-03-02", "period": 1, "studentId": "std_1", "status": "결석"}]}
        (OUT / "old.json").write_text(json.dumps(old, ensure_ascii=False), encoding="utf-8")
        await page.click('[data-target="data-manage"]')
        await page.set_input_files("#importFile", str(OUT / "old.json"))
        await page.click("#restore-btn")
        await page.wait_for_timeout(2500)
        await page.fill("#login-password", "after-recover")
        await page.click("#login-btn")
        await page.wait_for_selector("#login-overlay", state="hidden")
        check("v9 백업 복구(비밀번호 유지)", await page.evaluate("(async()=> (await StudentRepo.getAll())[0].name)()") == "옛학생")
        await page.evaluate("openAttendanceModalForDate('2026-03-02', 1, '국')")
        await page.wait_for_timeout(600)
        check("예전 기록 '미분류' 표시", await page.locator("#attendance-list-body .badge", has_text="미분류").count() == 1)
        await b.close()
    check("화면 오류 없음", not errors, errors[:5])


if __name__ == "__main__":
    port = free_port()
    data = tempfile.mkdtemp()
    threading.Thread(target=devserver.run, args=(port, data), daemon=True).start()
    import time
    time.sleep(1.0)
    asyncio.run(main(port))
    failed = [r for r in results if not r[1]]
    print(f"\n{len(results) - len(failed)}/{len(results)} 통과  (스크린샷·출력물: {OUT})")
    sys.exit(1 if failed else 0)
