import pytest

from backend.api import Api
from backend.db import Database
from backend import attendance as att
from tests.test_backend import setup_class


@pytest.fixture
def api(tmp_path):
    db = Database(tmp_path)
    db.unlock("1234")
    a = Api(db)
    setup_class(a, term1=("2026-03-02", "2026-03-13"), term2=("2026-09-01", "2026-09-04"))
    a.generate_annual()
    a.db_put_many("students", [{"studentId": "A", "number": 1, "name": "가"}, {"studentId": "B", "number": 2, "name": "나"}])
    return a


def rec(d, p, sid, status, reason=None, **kw):
    r = {"attendanceId": f"{d}_{p}_{sid}", "date": d, "period": p, "studentId": sid, "status": status}
    if reason:
        r["reason"] = reason
    r.update(kw)
    return r


def test_day_rules_neis():
    days = {"2026-03-02", "2026-03-03", "2026-03-04", "2026-03-05"}
    recs = [
        rec("2026-03-02", 1, "A", "결석", "질병"), rec("2026-03-02", 2, "A", "지각", "질병"),  # 결석일의 지각은 안 셈
        rec("2026-03-03", 1, "A", "지각", "미인정"), rec("2026-03-03", 2, "A", "조퇴", "기타"),
        rec("2026-03-04", 1, "A", "결석", "인정"), rec("2026-03-04", 2, "A", "결석", "인정"),
        rec("2026-03-05", 2, "A", "결과", "질병"), rec("2026-03-05", 3, "A", "결과", "질병"),
    ]
    c = att.student_counts(recs, days)
    assert c["결석_질병"] == 1 and c["지각_질병"] == 0
    assert c["지각_미인정"] == 1 and c["조퇴_기타"] == 1
    assert c["출석인정"] == 1 and c["결과_질병"] == 1
    assert c["수업일수"] == 4 and c["출석일수"] == 3
    old = att.student_counts([rec("2026-03-02", 1, "A", "결석")], days)
    assert old["결석_미분류"] == 1 and old["미분류"] == 1


def test_experiential_sync_approval_cycle(api):
    exp = {"studentId": "A", "studentName": "가", "type": "가족동반여행", "startDate": "2026-03-05", "endDate": "2026-03-09", "status": "requested"}
    r = api.save_experiential(exp)
    assert r["days"] == 0 and api.db_get_all("attendance") == []
    exp["status"] = "approved"
    r = api.save_experiential(exp)
    assert r["days"] == 3  # 목·금·월 (주말 제외)
    recs = api.db_get_all("attendance")
    assert len(recs) == 3 * 2 and all(x["reason"] == "인정" for x in recs)
    exp["endDate"] = "2026-03-05"
    assert api.save_experiential(exp)["days"] == 1
    assert len(api.db_get_all("attendance")) == 2
    api.delete_experiential(exp["expId"])
    assert api.db_get_all("attendance") == []


def test_regenerate_annual_resyncs_experiential(api):
    exp = {"studentId": "B", "studentName": "나", "type": "기타", "startDate": "2026-03-10", "endDate": "2026-03-10", "status": "approved"}
    api.save_experiential(exp)
    assert len(api.db_get_all("attendance")) == 2
    api.db_put("school_events", {"eventId": "h", "startDate": "2026-03-10", "endDate": "2026-03-10", "title": "재량휴업", "isHoliday": True})
    api.generate_annual()
    assert api.db_get_all("attendance") == []  # 휴업일이 되었으니 출석인정 기록도 사라짐


def test_neis_workbook(api):
    from openpyxl import load_workbook
    import base64, io
    api.db_put_many("attendance", [rec("2026-03-02", 1, "A", "결석", "질병"), rec("2026-03-02", 2, "A", "결석", "질병"),
                                   rec("2026-09-01", 1, "B", "지각", "미인정")])
    wb = load_workbook(io.BytesIO(base64.b64decode(api.build_file("neis_attendance_xlsx")["b64"])))
    assert wb.sheetnames[:3] == ["1학기", "2학기", "학년 전체"] and "3월" in wb.sheetnames and "9월" in wb.sheetnames
    ws = wb["1학기"]
    assert ws["D4"].value == "결석" and ws["D5"].value == "질병"
    assert ws["A6"].value == 1 and ws["C6"].value == 10 and ws["D6"].value == 1
    assert ws.cell(row=6, column=ws.max_column - 1).value == 9  # 출석일수
    assert wb["2학기"]["H7"].value == 1  # B: 지각(미인정)
    assert ws.print_title_rows == "$4:$5"
