import base64
import io
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.api import Api
from backend.db import Database, StoreError
from backend.excel_io import normalize_date

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def api(tmp_path):
    db = Database(tmp_path)
    assert db.unlock("1234")
    yield Api(db, ROOT / "curriculum_packs")
    db.close()


def xlsx_b64(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return base64.b64encode(bio.getvalue()).decode()


def setup_class(api, term1=("2026-03-02", "2026-03-31"), term2=("2026-09-01", "2026-09-30")):
    api.db_put("settings", {"id": "global", "schoolName": "테스트초", "schoolYear": "2026", "grade": "3", "classNo": "2",
                            "teacherName": "홍길동", "term1Start": term1[0], "term1End": term1[1],
                            "term2Start": term2[0], "term2End": term2[1]})
    api.db_put_many("subjects", [{"subjectId": "s1", "name": "수학", "shortName": "수", "hours": 10},
                                 {"subjectId": "s2", "name": "국어", "shortName": "국", "hours": 10}])
    grid = {}
    for d in ["mon", "tue", "wed", "thu", "fri"]:
        grid[f"{d}-1"] = {"id": "s2", "name": "국"}
        grid[f"{d}-2"] = {"id": "s1", "name": "수"}
    api.db_put("timetable_weekly", {"id": "weekly", "grid": grid})


# ---------------------------------------------------------------- 날짜
@pytest.mark.parametrize("raw,expected", [
    ("2026-03-02", "2026-03-02"), ("2026.3.2", "2026-03-02"), ("2026. 3. 2.", "2026-03-02"),
    ("2026/03/02", "2026-03-02"), ("3/2/2026", "2026-03-02"), ("2026년 3월 2일", "2026-03-02"),
    (46083, "2026-03-02"), (date(2026, 3, 2), "2026-03-02"), ("2026-03-02(월)", "2026-03-02"),
    ("abc", ""), ("", ""), (None, ""), ("2026-02-30", ""),
])
def test_normalize_date(raw, expected):
    assert normalize_date(raw) == expected


# ------------------------------------------------------------ 저장소
def test_crud_and_query(api):
    api.db_put("students", {"studentId": "a", "number": 1, "name": "가"})
    api.db_put_many("attendance", [
        {"attendanceId": "x1", "date": "2026-03-02", "period": 1, "studentId": "a", "status": "결석"},
        {"attendanceId": "x2", "date": "2026-03-03", "period": 1, "studentId": "a", "status": "출석"},
    ])
    assert api.db_get("students", "a")["name"] == "가"
    assert len(api.db_query("attendance", "studentId", "a")) == 2
    assert len(api.db_query("attendance", "date", "2026-03-02")) == 1
    api.delete_student_cascade("a")
    assert api.db_get("students", "a") is None
    assert api.db_get_all("attendance") == []


def test_missing_settings_returns_none_not_crash(api):
    assert api.db_get("settings", "global") is None
    with pytest.raises(StoreError, match="1학기 기간"):
        api.generate_annual()


def test_bad_store_rejected(api):
    with pytest.raises(StoreError):
        api.db_put("nope", {"id": 1})
    with pytest.raises(StoreError):
        api.db_put("students", {"name": "키 없음"})


# --------------------------------------------------------- 백업/복구
def test_import_old_v9_backup_keeps_current_password(api):
    old = {"settings": [{"id": "global", "schoolName": "옛학교"}, {"id": "security", "password": "9999"}],
           "students": [{"studentId": "std_1", "number": 1, "name": "옛학생"}],
           "attendance": [], "subjects": [], "custom_links": []}
    counts = api.import_backup(json.dumps(old))
    assert counts["students"] == 1 and counts["settings"] == 1
    assert api.verify_password("1234") and not api.verify_password("9999")
    assert api.db_get("settings", "security") is None


def test_backup_roundtrip(api):
    setup_class(api)
    f = api.build_file("backup")
    payload = base64.b64decode(f["b64"]).decode("utf-8")
    assert "수학" not in payload  # 백업 파일도 암호화
    api.clear_all()
    assert api.db_get_all("subjects") == []
    api.import_backup(payload)
    assert len(api.db_get_all("subjects")) == 2


def test_import_garbage_backup(api):
    with pytest.raises(StoreError):
        api.import_backup("not json")
    with pytest.raises(StoreError):
        api.import_backup(json.dumps({"hello": 1}))


def test_auto_backup_once_per_day(api):
    setup_class(api)
    db = api._db
    assert db.auto_backup_if_needed()
    assert db.auto_backup_if_needed() is None
    assert list((db.data_dir / "backups").glob("ilog_auto_*.json"))


# ---------------------------------------------------- 연간 시간표 생성
def test_generate_annual_keeps_content_and_skips_holidays(api):
    setup_class(api)
    api.db_put("school_events", {"eventId": "e1", "startDate": "2026-03-03", "endDate": "2026-03-03", "title": "개교기념일", "isHoliday": True})
    r = api.generate_annual()
    assert r["holidays"] == 1
    assert api.db_get("annual_schedule", "2026-03-03") is None
    assert api.db_get("annual_schedule", "2026-03-07") is None  # 토요일
    day = api.db_get("annual_schedule", "2026-03-02")
    day["subjects"][1]["content"] = "직접 입력한 내용"
    api.db_put("annual_schedule", day)
    r2 = api.generate_annual()
    assert r2["keptContent"] == 1
    assert api.db_get("annual_schedule", "2026-03-02")["subjects"][1]["content"] == "직접 입력한 내용"


# ------------------------------------------------------------ 엑셀 가져오기
def test_import_students_upsert_and_unique_ids(api):
    rows = [["번호", "이름", "성별", "전화번호"]] + [[i, f"학생{i}", "여" if i % 2 else "남", ""] for i in range(1, 31)]
    r = api.import_students(xlsx_b64(rows), "명단.xlsx")
    assert r["added"] == 30
    ids = {s["studentId"] for s in api.db_get_all("students")}
    assert len(ids) == 30  # 빠르게 넣어도 ID 중복 없음
    r2 = api.import_students(xlsx_b64([["번호", "이름"], [1, "개명학생"]]), "a.xlsx")
    assert r2 == {"added": 0, "updated": 1, "skipped": 0}
    assert len(api.db_get_all("students")) == 30


def test_import_students_requires_header(api):
    with pytest.raises(StoreError):
        api.import_students(xlsx_b64([["no", "nm"], [1, "a"]]), "a.xlsx")


def test_import_school_events_ranges_and_dupes(api):
    rows = [["날짜", "행사명", "휴업"], ["2026-07-20~2026-08-16", "여름방학", "O"], [date(2026, 3, 2), "입학식", ""],
            ["2026.5.5", "어린이날", "O"], ["엉터리", "x", ""]]
    r = api.import_school_events(xlsx_b64(rows), "학사.xlsx")
    assert r["added"] == 3 and r["badDates"] == ["엉터리"]
    ev = {e["title"]: e for e in api.db_get_all("school_events")}
    assert ev["여름방학"]["endDate"] == "2026-08-16" and ev["여름방학"]["isHoliday"]
    assert ev["어린이날"]["startDate"] == "2026-05-05"
    r2 = api.import_school_events(xlsx_b64(rows), "학사.xlsx")
    assert r2["added"] == 0 and r2["duplicates"] == 3


def test_import_subject_plan_with_varied_dates(api):
    setup_class(api)
    api.generate_annual()
    rows = [["날짜", "교시", "단원", "목표", "내용"], ["2026.3.2", 2, "1. 덧셈", "목표A", "내용A"],
            [date(2026, 3, 3), 2, "1. 덧셈", "목표B", ""], ["2026-03-02", 5, "x", "y", "z"]]
    r = api.import_subject_plan(xlsx_b64(rows), "p.xlsx", "수")
    assert r["applied"] == 2 and r["missingCount"] == 1
    s = api.db_get("annual_schedule", "2026-03-03")["subjects"][1]
    assert s["content"] == "1. 덧셈"  # 내용이 비면 단원으로


# ------------------------------------------------------------ 지도계획 라이브러리
def make_pack_rows(publisher="가나출판", n1=12, n2=5):
    rows = [["학년", "과목", "출판사", "학기", "차시", "단원", "학습목표", "지도내용", "교육과정"]]
    for i in range(1, n1 + 1):
        rows.append([3, "수학", publisher, 1, i, f"단원{(i - 1) // 4 + 1}", f"목표{i}", f"1학기 내용{i}", "2022 개정"])
    for i in range(1, n2 + 1):
        rows.append([3, "수학", publisher, 2, i, "2학기단원", f"목표{i}", f"2학기 내용{i}", "2022 개정"])
    return rows


def test_curriculum_import_preview_apply(api):
    setup_class(api, term1=("2026-03-02", "2026-03-13"), term2=("2026-09-01", "2026-09-04"))
    api.generate_annual()
    res = api.curriculum_import(xlsx_b64(make_pack_rows()), "pack.xlsx")
    assert res[0]["lessonCount"] == 17
    pid = res[0]["packId"]
    assert any(p["packId"] == pid for p in api.curriculum_list(3))
    pv = api.curriculum_preview(pid, "수")
    t1, t2 = pv["terms"]
    assert t1["slots"] == 10 and t1["lessons"] == 12 and t1["leftLessons"] == 2
    assert t2["slots"] == 4 and t2["lessons"] == 5
    r = api.curriculum_apply(pid, "수", True)
    assert r["applied"] == 14
    first = api.db_get("annual_schedule", "2026-03-02")["subjects"][1]
    assert first["content"] == "1학기 내용1" and first["lessonSeq"] == 1
    sep = api.db_get("annual_schedule", "2026-09-01")["subjects"][1]
    assert sep["content"] == "2학기 내용1"
    assert api.curriculum_choices()["수"] == pid
    # 국어 칸은 건드리지 않음
    assert api.db_get("annual_schedule", "2026-03-02")["subjects"][0]["content"] == ""


def test_curriculum_apply_keep_existing(api):
    setup_class(api, term1=("2026-03-02", "2026-03-06"), term2=("2026-09-01", "2026-09-01"))
    api.generate_annual()
    pid = api.curriculum_import(xlsx_b64(make_pack_rows()), "p.xlsx")[0]["packId"]
    d = api.db_get("annual_schedule", "2026-03-03")
    d["subjects"][1]["content"] = "내가 쓴 내용"
    api.db_put("annual_schedule", d)
    r = api.curriculum_apply(pid, "수", False)
    assert r["skipped"] == 1
    assert api.db_get("annual_schedule", "2026-03-03")["subjects"][1]["content"] == "내가 쓴 내용"
    assert api.db_get("annual_schedule", "2026-03-04")["subjects"][1]["content"] == "1학기 내용3"


def test_curriculum_seq_range_expansion_and_multi_pack(api):
    rows = [["학년", "과목", "출판사", "차시", "단원"], [4, "과학", "A사", "1~3", "물질"], [4, "과학", "B사", 1, "물질"]]
    res = api.curriculum_import(xlsx_b64(rows), "x.xlsx")
    counts = {p["publisher"]: p["lessonCount"] for p in res}
    assert counts == {"A사": 3, "B사": 1}


def test_curriculum_bad_header(api):
    with pytest.raises(StoreError, match="학년"):
        api.curriculum_import(xlsx_b64([["과목", "차시"], ["수학", 1]]), "x.xlsx")


def test_builtin_sample_pack_listed(api):
    packs = api.curriculum_list(3)
    assert any(p["builtin"] for p in packs)


# ------------------------------------------------------------------ 출력
def _load(b64):
    return load_workbook(io.BytesIO(base64.b64decode(b64)))


def populate(api):
    setup_class(api)
    api.generate_annual()
    api.import_students(xlsx_b64([["번호", "이름"], [1, "김<하늘>"], [2, "이바다"]]), "s.xlsx")
    sid = api.db_get_all("students")[0]["studentId"]
    api.db_put_many("attendance", [
        {"attendanceId": f"2026-03-02_{p}_{sid}", "date": "2026-03-02", "period": p, "studentId": sid, "status": "결석", "note": "감기"} for p in (1, 2)])
    api.db_put("eval_plans", {"planId": "p1", "subjectId": "s1", "subjectName": "수학", "title": "덧셈 평가", "date": "2026-03-10"})
    api.db_put("eval_scores", {"scoreId": f"p1_{sid}", "planId": "p1", "studentId": sid, "score": "상", "note": "잘함"})
    api.db_put("counseling", {"logId": "l1", "studentId": sid, "date": "2026-03-05", "type": "상담", "content": "교우관계"})
    api.db_put("incidents", {"incidentId": "i1", "studentId": sid, "date": "2026-03-06", "content": "비밀사안", "measures": ["상담"]})
    return sid


def test_xlsx_outputs(api):
    sid = populate(api)
    wb = _load(api.build_file("subject_plan_xlsx")["b64"])
    ws = wb["수학"]
    assert ws.page_setup.orientation == "landscape" and ws.print_title_rows == "$4:$4"
    assert ws["A4"].value == "차시" and ws["B5"].value == "2026-03-02"
    book = _load(api.build_file("class_book_xlsx")["b64"])
    assert "출결현황" in book.sheetnames and "사안기록(대외비)" not in book.sheetnames
    att = book["출결현황"]
    assert att["C5"].value == 1  # 결석 1일 (교시 2개여도 하루)
    assert "사안기록(대외비)" in _load(api.build_file("class_book_xlsx", {"incidents": True})["b64"]).sheetnames
    ev = api.build_file("eval_xlsx", {"planId": "p1"})
    assert ev["filename"].endswith(".xlsx")
    for k in ("template_subject_plan", "template_students", "template_school_events", "template_curriculum_pack"):
        assert _load(api.build_file(k)["b64"]).active.max_row >= 2


def test_html_outputs_escaped_and_complete(api):
    sid = populate(api)
    html = api.build_html("class_book")["html"]
    assert "김&lt;하늘&gt;" in html and "김<하늘>" not in html
    assert "thead { display: table-header-group; }" in html
    assert "비밀사안" not in html
    assert "비밀사안" in api.build_html("class_book", {"incidents": True})["html"]
    s = api.build_html("student_report", {"studentId": sid})["html"]
    assert "결석" in s and "교우관계" in s and "1,2" in s
    p = api.build_html("subject_plan")["html"]
    assert "수학 연간 지도 계획" in p


def test_empty_db_outputs_do_not_crash(api):
    for k in ("subject_plan_xlsx", "class_book_xlsx", "backup"):
        assert api.build_file(k)["b64"]
    for k in ("class_book", "subject_plan"):
        assert "<html" in api.build_html(k)["html"]
